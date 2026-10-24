"""generator.py - the Generator: pictures, presentations (.pptx) and narrated videos (.mp4) from a topic or from PDFs.

- Pictures: online image generation (Pollinations, free, no key) while the internet works; offline a painted
  placeholder (gradient + shapes + title) so a request never just fails.
- Presentations: the AI turns the topic / the PDFs' text into slides (title, bullets, speaker notes, a picture idea
  per slide); python-pptx builds a styled .pptx with a picture on every slide.
- Videos: the same outline becomes scenes: a picture with a slow zoom (Ken Burns), the caption on screen and the
  narration spoken (Kokoro online / Piper offline), joined into an .mp4 with moviepy.

Everything is saved in  Documents/Vermeil Creations/ . Long jobs run in the background and report progress
through on_event({"type": "gen", ...}).
"""
import hashlib
import io
import json
import math
import os
import random
import re
import textwrap
import threading
import time
import urllib.parse

OUT_DIR = os.path.join(os.path.expanduser("~"), "Documents", "Vermeil Creations")
on_event = None
_jobs = {}
THEMES = {
    "midnight": {"bg": (18, 24, 38), "bg2": (34, 48, 78), "title": (255, 214, 120), "text": (232, 236, 244), "accent": (224, 71, 91)},
    "parchment": {"bg": (243, 231, 200), "bg2": (222, 200, 150), "title": (110, 50, 20), "text": (50, 32, 18), "accent": (156, 107, 31)},
    "ocean": {"bg": (10, 52, 74), "bg2": (16, 104, 128), "title": (180, 245, 255), "text": (230, 248, 255), "accent": (255, 190, 90)},
    "forest": {"bg": (22, 46, 30), "bg2": (44, 88, 52), "title": (210, 255, 180), "text": (236, 246, 230), "accent": (240, 195, 90)},
    "clean": {"bg": (250, 250, 252), "bg2": (226, 232, 242), "title": (30, 40, 70), "text": (40, 46, 60), "accent": (63, 127, 214)},
}


def _emit(job, **kw):
    J = _jobs.get(job)
    if J is not None:
        J.update(kw)
    if on_event:
        try:
            on_event({"type": "gen", "job": job, **kw})
        except Exception:
            pass


def _safe(name):
    return re.sub(r"[^\w\- ]+", "", str(name or "")).strip()[:60] or time.strftime("Creation %Y-%m-%d %H%M")


def _out(name, ext):
    os.makedirs(OUT_DIR, exist_ok=True)
    base = os.path.join(OUT_DIR, _safe(name))
    p, n = base + ext, 2
    while os.path.exists(p):
        p, n = f"{base} ({n}){ext}", n + 1
    return p


# ---------------------------------------------------------------- reading PDFs
def pdf_text(paths, limit=60000):
    from pypdf import PdfReader
    out = []
    for p in paths or []:
        try:
            r = PdfReader(p)
            for page in r.pages[:80]:
                out.append(page.extract_text() or "")
                if sum(len(x) for x in out) > limit:
                    break
        except Exception as e:
            out.append(f"(couldn't read {os.path.basename(p)}: {e})")
    return re.sub(r"[ \t]+", " ", "\n".join(out))[:limit]


# ---------------------------------------------------------------- the outline (AI first, offline fallback)
def _sentences(text):
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if 25 < len(s.strip()) < 220]


def offline_outline(topic, source="", n=8):
    """no AI: slides from the PDF's own sentences (or a simple structure for a topic)."""
    title = (topic or "").strip() or (source.strip().splitlines() or ["Presentation"])[0][:60]
    slides = []
    if source:
        sents = _sentences(source)
        per = max(3, len(sents) // max(1, n))
        for i in range(0, min(len(sents), per * n), per):
            chunk = sents[i:i + per][:4]
            if not chunk: continue
            head = re.sub(r"[^\w\s]", "", chunk[0]).split()[:6]
            slides.append({"title": " ".join(head).title() or f"Part {len(slides) + 1}", "bullets": [c[:140] for c in chunk],
                           "notes": " ".join(chunk), "image": f"{title}, {' '.join(head)}"})
    else:
        for h in ["Introduction", "Background", "Key ideas", "How it works", "Examples", "Benefits", "Challenges", "Summary"][:n]:
            slides.append({"title": h, "bullets": [f"{h} of {title}"], "notes": f"{h} of {title}.", "image": f"{title} {h.lower()}"})
    return {"title": title, "subtitle": "Made with Vermeil", "slides": slides[:n] or [{"title": title, "bullets": [title], "notes": title, "image": title}]}


def outline(topic="", source="", n=8, kind="slides"):
    sys_p = ("You turn a topic or a document into a clear, well-ordered " + ("presentation" if kind == "slides" else "short narrated video") +
             f". Make exactly {n} {'slides' if kind == 'slides' else 'scenes'}. Facts must come from the document when one is given. "
             "Answer ONLY with JSON: {\"title\": ..., \"subtitle\": ..., \"slides\": [{\"title\": short heading, \"bullets\": [3-5 short points], "
             "\"notes\": speaker notes / narration, 2-4 sentences, \"image\": a vivid description of a picture for this slide (no text in it)}]}")
    user = f"TOPIC: {topic or '(from the document)'}\n" + (f"\nDOCUMENT:\n{source[:24000]}" if source else "")
    try:
        import llm
        r = llm.get_client().chat.completions.create(model="smart", temperature=0.4,
                                                     messages=[{"role": "system", "content": sys_p}, {"role": "user", "content": user}])
        m = re.search(r"\{.*\}", r.choices[0].message.content or "", re.S)
        if m:
            d = json.loads(m.group(0))
            slides = [{"title": str(s.get("title") or "")[:90], "bullets": [str(b)[:180] for b in (s.get("bullets") or [])][:6],
                       "notes": str(s.get("notes") or "")[:1200], "image": str(s.get("image") or s.get("title") or topic)[:300]}
                      for s in d.get("slides", []) if isinstance(s, dict)]
            if slides:
                return {"title": str(d.get("title") or topic or "Presentation")[:90], "subtitle": str(d.get("subtitle") or "")[:120], "slides": slides[:n], "ai": True}
    except Exception as e:
        print(f"[generator] AI outline failed ({e.__class__.__name__}) - using the offline outline.")
    return offline_outline(topic, source, n)


# ---------------------------------------------------------------- pictures
def _font(size, bold=False):
    from PIL import ImageFont
    for f in (["segoeuib.ttf", "arialbd.ttf"] if bold else ["segoeui.ttf", "arial.ttf"]) + ["DejaVuSans.ttf"]:
        try:
            return ImageFont.truetype(f, size)
        except OSError:
            continue
    return ImageFont.load_default()


def placeholder(prompt, w=1280, h=720, theme="midnight"):
    """offline picture: a painted gradient with soft shapes and the idea written on it."""
    from PIL import Image, ImageDraw, ImageFilter
    T = THEMES.get(theme, THEMES["midnight"])
    rnd = random.Random(hashlib.md5(prompt.encode()).hexdigest())
    img = Image.new("RGB", (w, h), T["bg"])
    d = ImageDraw.Draw(img)
    for y in range(h):
        k = y / h
        d.line([(0, y), (w, y)], fill=tuple(int(T["bg"][i] * (1 - k) + T["bg2"][i] * k) for i in range(3)))
    layer = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    dl = ImageDraw.Draw(layer)
    for _ in range(14):
        r = rnd.randint(h // 10, h // 2)
        x, y = rnd.randint(-r, w), rnd.randint(-r, h)
        c = tuple(min(255, int(v * rnd.uniform(0.8, 1.6))) for v in (T["accent"] if rnd.random() < 0.5 else T["bg2"])) + (rnd.randint(25, 70),)
        dl.ellipse([x, y, x + r * 2, y + r * 2], fill=c)
    layer = layer.filter(ImageFilter.GaussianBlur(28))
    img = Image.alpha_composite(img.convert("RGBA"), layer).convert("RGB")
    d = ImageDraw.Draw(img)
    f = _font(max(26, w // 26), bold=True)
    lines = textwrap.wrap(prompt, 36)[:4]
    y = h // 2 - len(lines) * f.size // 1.6
    for line in lines:
        tw = d.textlength(line, font=f)
        d.text(((w - tw) / 2, y), line, font=f, fill=T["title"], stroke_width=2, stroke_fill=(0, 0, 0))
        y += f.size * 1.25
    return img


_last_free = [0.0]              # when the last free (keyless) picture was asked for


def _key():
    return (os.environ.get("POLLINATIONS_TOKEN") or os.environ.get("POLLINATIONS_KEY") or "").strip()


def picture(prompt, w=1280, h=720, seed=None, theme="midnight", online=True, job=None, tries=4):
    """-> PIL image, painted by the online image AI (Pollinations).
    Without a key it allows 1 picture every 45 s per computer, so we wait our turn (with a countdown) and retry
    instead of giving up. With POLLINATIONS_TOKEN in the env file there's no waiting. Only if the internet is down
    (or it keeps refusing) do you get the offline painted card."""
    from PIL import Image
    if online:
        import requests
        key = _key()
        for attempt in range(tries):
            if not key:                                         # space keyless requests 46 s apart
                wait = 46 - (time.time() - _last_free[0])
                if wait > 0: _countdown(job, wait, "The free image AI paints 1 picture every 45 s")
                _last_free[0] = time.time()
            try:
                url = ("https://image.pollinations.ai/prompt/" + urllib.parse.quote(prompt[:400]) +
                       f"?width={w}&height={h}&nologo=true&model=flux&seed={seed if seed is not None else random.randint(1, 10 ** 6)}")
                r = requests.get(url, timeout=150, headers={"Authorization": f"Bearer {key}"} if key else {})
                if r.ok and r.headers.get("content-type", "").startswith("image"):
                    return Image.open(io.BytesIO(r.content)).convert("RGB").resize((w, h))
                m = re.search(r"Retry after (\d+)", r.text or "")
                if r.status_code in (402, 429) or m:
                    _last_free[0] = time.time() - 46 + (int(m.group(1)) + 1 if m else 46)
                    continue
                print(f"[generator] image AI answered {r.status_code}: {(r.text or '')[:120]}")
            except Exception as e:
                print(f"[generator] online picture failed ({e.__class__.__name__}) - trying again.")
                time.sleep(3)
    return placeholder(prompt, w, h, theme)


def _countdown(job, seconds, why):
    end = time.time() + seconds
    while time.time() < end:
        _emit(job, text=f"{why} - next one in {int(end - time.time()) + 1} s…")
        time.sleep(min(1.0, max(0.05, end - time.time())))


def pictures(prompt, count=4, size="landscape", style="", job=None):
    w, h = {"landscape": (1280, 720), "portrait": (720, 1280), "square": (1024, 1024)}.get(size, (1280, 720))
    full = f"{prompt}, {style}" if style else prompt
    paths = []
    for i in range(max(1, min(8, int(count)))):
        _emit(job, progress=i / count, text=f"Painting picture {i + 1} of {count}…")
        img = picture(full, w, h, seed=random.randint(1, 10 ** 6), job=job)
        p = _out(f"{prompt[:40]} {i + 1}", ".png")
        img.save(p)
        paths.append(p)
    return {"files": paths}


# ---------------------------------------------------------------- presentations
def presentation(topic="", pdfs=None, slides=8, theme="midnight", pictures_on=True, job=None):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    T = THEMES.get(theme, THEMES["midnight"])
    _emit(job, progress=0.02, text="Reading your PDFs…" if pdfs else "Planning the slides…")
    src = pdf_text(pdfs) if pdfs else ""
    O = outline(topic, src, slides, "slides")
    _emit(job, progress=0.15, text=f"Writing {len(O['slides'])} slides…")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    col = lambda c: RGBColor(*c)

    def background(s):
        f = s.background.fill
        f.gradient(); f.gradient_angle = 90
        f.gradient_stops[0].color.rgb = col(T["bg"]); f.gradient_stops[1].color.rgb = col(T["bg2"])

    def text(s, x, y, w, h, txt, size, color, bold=False, bullets=None):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        items = bullets if bullets is not None else [txt]
        for i, line in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = ("•  " if bullets is not None else "") + line
            p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = col(color)
            p.space_after = Pt(10 if bullets is not None else 0)
        return tb

    # title slide
    s = prs.slides.add_slide(blank); background(s)
    if pictures_on:
        img = picture(f"{O['title']}, cinematic, beautiful", 1280, 720, theme=theme, job=job)
        bio = io.BytesIO(); img.save(bio, "PNG"); bio.seek(0)
        s.shapes.add_picture(bio, 0, 0, prs.slide_width, prs.slide_height)
        shade = s.shapes.add_shape(1, 0, Inches(4.3), prs.slide_width, Inches(3.2))
        shade.fill.solid(); shade.fill.fore_color.rgb = col(T["bg"]); shade.line.fill.background()
    text(s, 0.8, 4.6, 11.7, 1.4, O["title"], 44, T["title"], bold=True)
    text(s, 0.8, 5.9, 11.7, 0.9, O.get("subtitle") or "", 22, T["text"])
    previews = []
    for i, sl in enumerate(O["slides"]):
        _emit(job, progress=0.15 + 0.8 * i / len(O["slides"]), text=f"Slide {i + 2}: {sl['title']}")
        s = prs.slides.add_slide(blank); background(s)
        bar = s.shapes.add_shape(1, Inches(0.6), Inches(1.35), Inches(1.2), Inches(0.08))
        bar.fill.solid(); bar.fill.fore_color.rgb = col(T["accent"]); bar.line.fill.background()
        text(s, 0.6, 0.45, 12, 0.9, sl["title"], 34, T["title"], bold=True)
        tw = 6.6 if pictures_on else 12
        text(s, 0.6, 1.7, tw, 5.2, "", 20, T["text"], bullets=sl["bullets"] or [sl["notes"][:160]])
        if pictures_on:
            img = picture(sl["image"], 1024, 1024, theme=theme, job=job)
            bio = io.BytesIO(); img.save(bio, "PNG"); bio.seek(0)
            s.shapes.add_picture(bio, Inches(7.6), Inches(1.5), Inches(5.2), Inches(5.2))
        s.notes_slide.notes_text_frame.text = sl["notes"]
        previews.append({"title": sl["title"], "bullets": sl["bullets"]})
    path = _out(O["title"], ".pptx")
    prs.save(path)
    return {"file": path, "title": O["title"], "slides": previews, "ai": bool(O.get("ai"))}


# ---------------------------------------------------------------- videos
def _narrate(text):
    """(float32 samples, rate) - the same voice rules as everywhere: Kokoro while online, Piper offline."""
    import numpy as np
    try:
        import pinakamain as pm
        import tts
        v = pm.active_voice()
        samples, rate = tts.synth(text, v, "+0%", "+0Hz")
        return np.asarray(samples, dtype=np.float32).reshape(-1), int(rate)
    except Exception as e:
        print(f"[generator] narration failed ({e.__class__.__name__}) - silent scene.")
        return np.zeros(int(22050 * max(3, len(text) / 15)), dtype=np.float32), 22050


def _frame_with_caption(img, caption, W, H, T):
    from PIL import ImageDraw
    im = img.copy()
    d = ImageDraw.Draw(im, "RGBA")
    f = _font(max(22, W // 46), bold=True)
    lines = textwrap.wrap(caption, 52)[:3]
    if lines:
        bh = int(len(lines) * f.size * 1.35 + 30)
        d.rectangle([0, H - bh, W, H], fill=(0, 0, 0, 150))
        y = H - bh + 15
        for line in lines:
            tw = d.textlength(line, font=f)
            d.text(((W - tw) / 2, y), line, font=f, fill=(255, 255, 255), stroke_width=2, stroke_fill=(0, 0, 0))
            y += f.size * 1.35
    return im


def video(topic="", pdfs=None, scenes=6, theme="midnight", job=None, W=1280, H=720, fps=24):
    import numpy as np
    import soundfile as sf
    try:                                               # moviepy 2.x
        from moviepy import AudioFileClip, VideoClip, concatenate_videoclips
    except ImportError:                                # moviepy 1.x
        from moviepy.editor import AudioFileClip, VideoClip, concatenate_videoclips
    with_audio = lambda c, a: c.with_audio(a) if hasattr(c, "with_audio") else c.set_audio(a)
    T = THEMES.get(theme, THEMES["midnight"])
    _emit(job, progress=0.02, text="Reading your PDFs…" if pdfs else "Writing the script…")
    src = pdf_text(pdfs) if pdfs else ""
    O = outline(topic, src, scenes, "video")
    tmp = os.path.join(OUT_DIR, ".tmp"); os.makedirs(tmp, exist_ok=True)
    clips = []
    allscenes = [{"title": O["title"], "notes": O["title"] + (". " + O["subtitle"] if O.get("subtitle") else ""), "image": f"{O['title']}, cinematic"}] + O["slides"]
    for i, sc in enumerate(allscenes):
        _emit(job, progress=0.05 + 0.8 * i / len(allscenes), text=f"Scene {i + 1} of {len(allscenes)}: {sc['title']}")
        img = picture(sc["image"], W, H, theme=theme, job=job)
        narr = sc["notes"] or sc["title"]
        samples, rate = _narrate(narr)
        wav = os.path.join(tmp, f"s{i}.wav"); sf.write(wav, samples, rate)
        dur = max(3.0, len(samples) / rate + 0.6)
        base = _frame_with_caption(img, sc["title"] if i == 0 else narr[:150], W, H, T)
        arr = np.asarray(base)
        zoom_in = i % 2 == 0

        def make_frame(t, arr=arr, dur=dur, zoom_in=zoom_in):
            k = t / dur
            z = 1 + 0.08 * (k if zoom_in else 1 - k)                  # slow Ken Burns zoom
            h, w = arr.shape[:2]
            ch, cw = int(h / z), int(w / z)
            y0, x0 = (h - ch) // 2, (w - cw) // 2
            crop = arr[y0:y0 + ch, x0:x0 + cw]
            import cv2
            return cv2.resize(crop, (w, h), interpolation=cv2.INTER_LINEAR)

        clip = with_audio(VideoClip(make_frame, duration=dur), AudioFileClip(wav))
        clips.append(clip)
    _emit(job, progress=0.88, text="Putting the video together (this takes a moment)…")
    final = concatenate_videoclips(clips, method="compose")
    path = _out(O["title"], ".mp4")
    final.write_videofile(path, fps=fps, codec="libx264", audio_codec="aac", logger=None, threads=4)
    for c in clips:
        try: c.close()
        except Exception: pass
    return {"file": path, "title": O["title"], "scenes": len(allscenes), "seconds": round(final.duration, 1), "ai": bool(O.get("ai"))}


# ---------------------------------------------------------------- background jobs
def start(kind, **kw):
    """start a job in the background -> {'job': id}. Progress and the result arrive as 'gen' events."""
    fn = {"pictures": pictures, "presentation": presentation, "video": video}.get(kind)
    if not fn:
        return {"error": "Unknown kind - pick pictures, presentation or video."}
    if not (kw.get("topic") or kw.get("prompt") or kw.get("pdfs")):
        return {"error": "Type a topic (or a picture idea), or add PDFs."}
    job = f"g{int(time.time() * 1000) % 10 ** 9}"
    _jobs[job] = {"kind": kind, "state": "running", "progress": 0, "text": "Starting…", "started": time.time()}

    def run():
        try:
            res = fn(job=job, **kw)
            _emit(job, state="done", progress=1, text="Done!", result=res)
        except Exception as e:
            import traceback; traceback.print_exc()
            _emit(job, state="error", text=f"{e.__class__.__name__}: {e}")

    threading.Thread(target=run, daemon=True, name=f"gen-{kind}").start()
    return {"job": job}


def status(job=None):
    if job:
        return _jobs.get(job) or {"error": "No such job."}
    return {"jobs": _jobs, "folder": OUT_DIR}


def recent(n=30):
    if not os.path.isdir(OUT_DIR):
        return []
    files = [os.path.join(OUT_DIR, f) for f in os.listdir(OUT_DIR) if f.lower().endswith((".png", ".pptx", ".mp4"))]
    files.sort(key=os.path.getmtime, reverse=True)
    return [{"path": p, "name": os.path.basename(p), "kind": os.path.splitext(p)[1][1:], "size_kb": os.path.getsize(p) // 1024,
             "saved": time.strftime("%Y-%m-%d %H:%M", time.localtime(os.path.getmtime(p)))} for p in files[:n]]
