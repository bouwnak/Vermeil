"""
gym.py - Gym mode: your workout routine, a log of every set (day by day), body stats, and progress tips.

Files (in data/):
  gym_routine.json - the routine (imported from your HYBRID.xlsx, or edited in the app)
  gym_log.json     - {"2026-10-20": {"day": 1, "title": "PUSH", "sets": [{"exercise", "set", "kg", "reps", "time"}]}}
  gym_stats.json   - [{"date", "weight", "biceps_r", ...}]

Import your Excel file:  python gym.py "C:\\path\\to\\HYBRID.xlsx"
"""
import os as _os, sys as _sys   # lets this file run on its own too (VS Code "Run"): find brain/, senses/, skills/
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
import folders  # noqa: E402,F401
import json
import os
import re
import sys
from datetime import date, datetime, timedelta

from paths import data_file

ROUTINE_FILE = data_file("gym_routine.json")
LOG_FILE = data_file("gym_log.json")
STATS_FILE = data_file("gym_stats.json")
STAT_FIELDS = ["weight", "biceps_r", "biceps_l", "thigh_r", "thigh_l", "chest", "waist"]
DEFAULT_REST = 90


def _read(path, default):
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError):
        return default


def _write(path, data):
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    os.replace(tmp, path)


# ------------------------------------------------------------------ importing the Excel file
def _reps_text(v):
    """Excel turned '6-12' into the date 12 June. Undo that."""
    if isinstance(v, (datetime, date)):
        return f"{v.month}-{v.day}"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip() if v is not None else ""


def _num(v):
    if isinstance(v, (int, float)):
        return float(v)
    m = re.match(r"\s*([\d.]+)", str(v or ""))
    return float(m.group(1)) if m else None


def _rest_seconds(text):
    nums = [int(n) for n in re.findall(r"\d+", str(text or ""))]
    if not nums:
        return DEFAULT_REST
    secs = max(nums) * (60 if "min" in str(text).lower() else 1)
    return secs


def _is_link(v):
    return isinstance(v, str) and v.strip().lower().startswith("http")


def parse_workbook(path):
    """Read a routine sheet laid out like HYBRID.xlsx: blocks of 3 days side by side,
    columns: name, rest, set#, completed, kg, reps needed, reps done, link."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    routine = {"name": os.path.splitext(os.path.basename(path))[0], "days": [], "profile": {}}
    sheet = next((ws for ws in wb if ws.title.upper() != "STATISTICS"), wb.active)
    rows = [[c for c in r] for r in sheet.iter_rows(values_only=True)]
    width = max(len(r) for r in rows) if rows else 0
    rows = [list(r) + [None] * (width - len(r)) for r in rows]
    cell = lambda r, c: rows[r][c] if 0 <= r < len(rows) and 0 <= c < width else None

    # profile (name, calories...) - label in one cell, value in the next filled cell
    for r in range(min(20, len(rows))):
        for c in range(width):
            label = cell(r, c)
            if isinstance(label, str) and label.strip().endswith(":"):
                val = next((cell(r, k) for k in range(c + 1, min(c + 3, width)) if cell(r, k) is not None), None)
                if val is not None:
                    routine["profile"].setdefault(label.strip(" :").lower(), val)

    # day headers look like "1st day" with the title in the next column
    headers = [(r, c) for r in range(len(rows)) for c in range(width)
               if isinstance(cell(r, c), str) and re.fullmatch(r"\s*\d+(st|nd|rd|th) day\s*", cell(r, c), re.I)]
    for i, (r, c) in enumerate(headers):
        num = int(re.match(r"\s*(\d+)", cell(r, c)).group(1))
        title = str(cell(r, c + 1) or f"Day {num}").strip()
        end = next((r2 for r2, c2 in headers[i + 1:] if r2 > r), len(rows))
        has_table = any(cell(k, c) == "Working Sets" for k in range(r, min(r + 4, end)))
        day = {"day": num, "title": title, "exercises": []}
        if not has_table:                                   # e.g. cardio day: name | time | notes
            for k in range(r + 1, end):
                name = cell(k, c)
                if isinstance(name, str) and name.strip() and not _is_link(name):
                    if name.strip().upper().startswith("SETTING UP"):
                        break
                    day["exercises"].append({"name": name.strip(), "sets": 1, "target": _reps_text(cell(k, c + 1)),
                                             "kg": None, "rest": 0, "notes": str(cell(k, c + 3) or "").strip(),
                                             "link": "", "timed": True})
            routine["days"].append(day)
            continue
        ex = None
        for k in range(r + 2, end):
            name, rest, setno, kg, target = cell(k, c), cell(k, c + 1), cell(k, c + 2), cell(k, c + 4), cell(k, c + 5)
            link = next((cell(k, c + j) for j in (6, 7) if _is_link(cell(k, c + j))), "")
            if setno == 1 and isinstance(name, str) and name.strip() and not _is_link(name):
                t = _reps_text(target)
                ex = {"name": name.strip(), "sets": 1, "target": t, "kg": _num(kg), "kg_text": "" if _num(kg) is not None or kg is None else str(kg),
                      "rest": _rest_seconds(rest), "notes": "", "link": link or "",
                      "timed": bool(re.search(r"sec|min", t, re.I))}
                day["exercises"].append(ex)
                continue
            if ex is None:
                continue
            if isinstance(setno, (int, float)) and setno > 1:
                ex["sets"] = max(ex["sets"], int(setno))
                if _num(kg) is not None:
                    ex["kg"] = max(ex["kg"] or 0, _num(kg))
            if isinstance(name, str) and name.strip() and not _is_link(name):
                ex["notes"] = (ex["notes"] + " " + name.strip()).strip()
            if link and not ex["link"]:
                ex["link"] = link
        routine["days"].append(day)
    for d in routine["days"]:
        for ex in d["exercises"]:
            if not ex["target"] and "AMRAP" in ex["notes"].upper():
                ex["target"] = "AMRAP"
    routine["days"].sort(key=lambda d: d["day"])

    # STATISTICS sheet -> body stats
    stats = []
    if "STATISTICS" in [ws.title.upper() for ws in wb]:
        ws = next(ws for ws in wb if ws.title.upper() == "STATISTICS")
        for row in ws.iter_rows(values_only=True):
            row = [v for v in row]
            first = next((i for i, v in enumerate(row) if isinstance(v, str) and v.strip().lower().startswith("week")), None)
            if first is None:
                continue
            vals = [_num(v) for v in row[first + 1:first + 1 + len(STAT_FIELDS)]]
            if any(v is not None for v in vals):
                stats.append({"date": row[first].strip(), **{f: v for f, v in zip(STAT_FIELDS, vals)}})
    return routine, stats


def import_xlsx(path):
    routine, stats = parse_workbook(path)
    if not routine["days"]:
        return {"error": "I couldn't find any days like '1st day' in that file."}
    _write(ROUTINE_FILE, routine)
    if stats:
        old = [s for s in _read(STATS_FILE, []) if not str(s.get("date", "")).lower().startswith("week")]
        _write(STATS_FILE, stats + old)
    n = sum(len(d["exercises"]) for d in routine["days"])
    return {"message": f"Imported {len(routine['days'])} days and {n} exercises.", "days": len(routine["days"])}


# ------------------------------------------------------------------ routine & log
def get_routine():
    return _read(ROUTINE_FILE, {"name": "", "days": [], "profile": {}})


def _clean_exercise(ex):
    target = str(ex.get("target") or "").strip()
    return {"name": str(ex.get("name") or "New exercise").strip()[:80],
            "sets": max(1, min(20, int(_num(ex.get("sets")) or 1))),
            "target": target[:30],
            "kg": _num(ex.get("kg")) if ex.get("kg") not in (None, "") else None,
            "kg_text": str(ex.get("kg_text") or "")[:20],
            "rest": max(0, min(900, int(_num(ex.get("rest")) or 0))),
            "notes": str(ex.get("notes") or "")[:300],
            "link": str(ex.get("link") or "").strip() if _is_link(ex.get("link")) else "",
            "timed": bool(ex.get("timed")) if "timed" in ex else bool(re.search(r"sec|min", target, re.I))}


def save_routine(routine):
    """Save an edited routine (from the Gym editor). Days are renumbered 1, 2, 3... in order."""
    if not isinstance(routine, dict) or not isinstance(routine.get("days"), list):
        return {"error": "Bad routine."}
    old = get_routine()
    days = []
    for i, d in enumerate(routine["days"], 1):
        if not isinstance(d, dict):
            continue
        days.append({"day": i, "title": str(d.get("title") or f"Day {i}").strip()[:40],
                     "exercises": [_clean_exercise(e) for e in d.get("exercises", []) if isinstance(e, dict)]})
    _write(ROUTINE_FILE, {"name": routine.get("name", old.get("name", "")), "profile": old.get("profile", {}), "days": days})
    return {"message": "Routine saved.", "days": len(days)}


def upgrade_weight(exercise, new_kg=None, step=2.5):
    """Raise the planned weight of an exercise (everywhere it appears in the routine). Returns the new weight."""
    r = get_routine()
    found = [ex for d in r["days"] for ex in d["exercises"] if ex["name"].lower() == exercise.lower()]
    if not found:
        return {"error": f"{exercise} isn't in your routine."}
    if new_kg is None:
        last = last_performance(exercise)
        base = max([ex.get("kg") or 0 for ex in found] + [s.get("kg") or 0 for s in (last or {}).get("sets", [])])
        new_kg = base + step
    new_kg = round(float(new_kg) * 4) / 4                  # keep to 0.25 kg steps
    for ex in found:
        ex["kg"] = new_kg
    save_routine(r)
    return {"kg": new_kg, "message": f"{exercise}: {new_kg:g} kg from now on."}


def new_day(title="New day"):
    r = get_routine()
    r["days"].append({"title": title, "exercises": []})
    return save_routine(r)


def reset_day(when=None):
    """Restart training: clear the ticks of the session you're doing NOW and start it again.
    Finished sessions, History, bests and body stats are never touched."""
    when = when or date.today().isoformat()
    log = _log()
    entry = log.get(when)
    current = _session(entry)
    cleared = len(_session_sets(entry, current))
    if cleared:
        entry["sets"] = [st for st in entry["sets"] if int(st.get("session", 1)) != current]
        entry.get("session_info", {}).pop(str(current), None)
        _write(LOG_FILE, log)
    return {"message": f"Restarted this session - cleared {cleared} ticked set{'s' if cleared != 1 else ''}. "
                       f"Finished sessions are still in History." if cleared else "This session is already empty.",
            "cleared": cleared, "session": current}


def _session(entry):
    return int((entry or {}).get("session", 1))


def _session_sets(entry, session):
    return [st for st in (entry or {}).get("sets", []) if int(st.get("session", 1)) == session]


def _is_current(d, st, log, today=None):
    """Is this set part of the session you're doing right now?"""
    today = today or date.today().isoformat()
    return d == today and int(st.get("session", 1)) == _session(log.get(today))


def _log():
    data = _read(LOG_FILE, {})
    return data if isinstance(data, dict) else {}


BACKUP_FILE = data_file("gym_log_backup.json")


def clear_history():
    """Clear every logged workout (History starts fresh). A copy is kept so 'Undo' can bring it back.
    Your routine, weights and body stats are NOT touched."""
    log = _log()
    if not log:
        return {"cleared": 0}
    _write(BACKUP_FILE, log)
    _write(LOG_FILE, {})
    return {"cleared": len(log)}


def undo_clear_history():
    old = _read(BACKUP_FILE, None)
    if not isinstance(old, dict) or not old:
        return {"restored": 0}
    log = _log()
    for day, entry in old.items():                 # merge, so anything logged after clearing is kept too
        log.setdefault(day, entry)
    _write(LOG_FILE, log)
    os.remove(BACKUP_FILE)
    return {"restored": len(old)}


def next_day(today=None):
    """Which routine day comes next: the one after the last day you trained (wraps around)."""
    days = [d["day"] for d in get_routine()["days"]]
    if not days:
        return None
    today = (today or date.today()).isoformat()
    log = _log()
    cur = _session_sets(log.get(today), _session(log.get(today)))
    if cur:                                             # in the middle of a session: stay on that day
        d = next((s.get("day") for s in cur if s.get("day")), log[today].get("day"))
        if d in days:
            return d
    past = sorted(k for k in log if k <= today and log[k].get("sets"))
    if not past:
        return days[0]
    e = log[past[-1]]
    last_session = max(int(s.get("session", 1)) for s in e["sets"])
    last = next((s.get("day") for s in e["sets"] if int(s.get("session", 1)) == last_session and s.get("day")), e.get("day"))
    return days[(days.index(last) + 1) % len(days)] if last in days else days[0]


def day_plan(day_num=None):
    routine = get_routine()
    day_num = day_num or next_day()
    day = next((d for d in routine["days"] if d["day"] == day_num), None)
    if not day:
        return None
    out = json.loads(json.dumps(day))
    for ex in out["exercises"]:
        ex["last"] = last_performance(ex["name"])
        ex["best"] = {str(k): v for k, v in best_sets(ex["name"]).items()}      # str keys: JSON-safe
        ex["tip"] = overload_tip(ex)
    return out


def log_set(exercise, set_no, kg=None, reps=None, day=None, when=None):
    when = when or date.today().isoformat()
    log = _log()
    entry = log.setdefault(when, {"day": day, "title": "", "sets": []})
    session = _session(entry)
    title = entry.get("title", "")
    if day:
        entry["day"] = day
        d = next((d for d in get_routine()["days"] if d["day"] == day), None)
        title = entry["title"] = d["title"] if d else title
    entry["sets"] = [s for s in entry["sets"] if not (s["exercise"] == exercise and s["set"] == int(set_no)
                                                      and int(s.get("session", 1)) == session)]
    entry["sets"].append({"exercise": exercise, "set": int(set_no), "kg": _num(kg), "reps": str(reps or "").strip(),
                          "time": datetime.now().strftime("%H:%M"), "session": session, "day": day or entry.get("day"),
                          "title": title})
    _write(LOG_FILE, log)
    return entry


def unlog_set(exercise, set_no, when=None):
    when = when or date.today().isoformat()
    log = _log()
    if when in log:
        session = _session(log[when])
        log[when]["sets"] = [s for s in log[when]["sets"] if not (s["exercise"] == exercise and s["set"] == int(set_no)
                                                                  and int(s.get("session", 1)) == session)]
        _write(LOG_FILE, log)
    return True


def finish_workout(minutes, note="", when=None):
    """Finish: save this session (its sets + how long it took) to History, then start a new, empty session."""
    when = when or date.today().isoformat()
    log = _log()
    entry = log.get(when)
    current = _session(entry)
    sets = _session_sets(entry, current)
    if not sets:
        return {"saved": False, "message": "Nothing ticked in this session, so there's nothing to save."}
    entry.setdefault("session_info", {})[str(current)] = {"minutes": round(float(minutes or 0)), "note": note or "",
                                                          "finished": datetime.now().strftime("%H:%M")}
    entry["session"] = current + 1
    _write(LOG_FILE, log)
    return {"saved": True, "session": current, "sets": len(sets), "minutes": round(float(minutes or 0)),
            "message": f"Session saved: {round(float(minutes or 0))} min, {len(sets)} set{'s' if len(sets) != 1 else ''}. "
                       f"Nice work! The next session starts fresh."}


def today_log(when=None):
    """Only the sets of the session you're doing now (earlier sessions today are in History)."""
    entry = _log().get(when or date.today().isoformat())
    if not entry:
        return {"day": None, "sets": [], "session": 1}
    session = _session(entry)
    return {**entry, "sets": _session_sets(entry, session), "session": session}


def last_performance(exercise, before=None):
    """The most recent earlier session with this exercise (an earlier session today counts too)."""
    today = date.today().isoformat()
    before = before or today
    log = _log()
    for d in sorted(log, reverse=True):
        if d > before:
            continue
        mine = [s for s in log[d]["sets"] if s["exercise"].lower() == exercise.lower()
                and not (d == before == today and _is_current(d, s, log, today)) and not (d == before != today)]
        if not mine:
            continue
        last_session = max(int(s.get("session", 1)) for s in mine)
        sets = sorted((s for s in mine if int(s.get("session", 1)) == last_session), key=lambda s: s["set"])
        return {"date": d, "sets": sets}
    return None


def rep_value(reps):
    """How much a 'reps done' entry is worth, for comparing: '12' -> 12, '30 sec' -> 30, '1min 10 secs' -> 70."""
    text = str(reps or "").lower()
    mins = re.search(r"(\d+(?:\.\d+)?)\s*m(?:in)?", text)
    secs = re.search(r"(\d+(?:\.\d+)?)\s*s(?:ec)?", text)
    if mins or secs:
        return (float(mins.group(1)) * 60 if mins else 0) + (float(secs.group(1)) if secs else 0)
    return _num(text)


def best_sets(exercise, before=None):
    """Your record for each set number of this exercise, from every EARLIER session (not the one you're doing now):
    {set_no: {'reps': '12', 'kg': 20.0, 'date': '2026-10-14'}}. Most reps wins; same reps -> heavier wins."""
    today = date.today().isoformat()
    before = before or today
    log = _log()
    best = {}
    for d, entry in log.items():
        if d > before or (d == before != today):
            continue
        for st in entry.get("sets", []):
            if st["exercise"].lower() != exercise.lower() or (d == before == today and _is_current(d, st, log, today)):
                continue
            v = rep_value(st.get("reps"))
            if v is None:
                continue
            cur = best.get(st["set"])
            key = (v, st.get("kg") or 0)
            if cur is None or key > (cur["_v"], cur["kg"] or 0):
                best[st["set"]] = {"reps": st.get("reps"), "kg": st.get("kg"), "date": d, "_v": v}
    return {n: {k: v for k, v in b.items() if k != "_v"} | {"value": b["_v"]} for n, b in best.items()}


def _target_top(target):
    nums = [int(n) for n in re.findall(r"\d+", target or "")]
    return max(nums) if nums else None


def overload_tip(ex):
    """Progressive overload: hit the top of the rep range on every set -> add weight."""
    last = ex.get("last") or last_performance(ex["name"])
    top = _target_top(ex.get("target"))
    if not last or ex.get("timed") or not top:
        return ""
    reps = [_num(s["reps"]) for s in last["sets"]]
    kg = max((s["kg"] or 0) for s in last["sets"])
    if len(reps) >= ex.get("sets", 1) and all(r is not None and r >= top for r in reps):
        step = 2.5 if kg >= 10 else 1.25
        return f"You hit {top} reps on every set last time - try {kg + step:g} kg."
    if kg:
        return f"Last time: {kg:g} kg. Beat your reps today."
    return ""


def history(days=30):
    """One row per session (a day can have several), newest first."""
    log = _log()
    cutoff = (date.today() - timedelta(days=days)).isoformat()
    out = []
    for d in sorted(log, reverse=True):
        e = log[d]
        if d < cutoff or not e.get("sets"):
            continue
        sessions = sorted({int(s.get("session", 1)) for s in e["sets"]}, reverse=True)
        for n in sessions:
            sets = _session_sets(e, n)
            info = e.get("session_info", {}).get(str(n), {})
            volume = sum((s["kg"] or 0) * (_num(s["reps"]) or 0) for s in sets)
            title = next((s.get("title") for s in sets if s.get("title")), e.get("title", ""))
            out.append({"date": d, "session": n, "sessions_that_day": len(sessions),
                        "day": next((s.get("day") for s in sets if s.get("day")), e.get("day")), "title": title,
                        "sets": len(sets), "exercises": len({s["exercise"] for s in sets}), "volume": round(volume),
                        "minutes": info.get("minutes", e.get("minutes") if len(sessions) == 1 else None),
                        "current": d == date.today().isoformat() and n == _session(e), "detail": sets})
    return out


def exercise_progress(exercise):
    """Best weight per day for one exercise, oldest first (for a chart)."""
    log = _log()
    points = []
    for d in sorted(log):
        sets = [s for s in log[d]["sets"] if s["exercise"].lower() == exercise.lower() and s["kg"]]
        if sets:
            points.append({"date": d, "kg": max(s["kg"] for s in sets)})
    return points


def streak(today=None):
    today = today or date.today()
    trained = {d for d, e in _log().items() if e.get("sets")}
    n, d = 0, today
    if d.isoformat() not in trained:
        d -= timedelta(days=1)
    while d.isoformat() in trained:
        n += 1
        d -= timedelta(days=1)
    return n


# ------------------------------------------------------------------ body stats
def get_stats():
    data = _read(STATS_FILE, [])
    return data if isinstance(data, list) else []


def add_stats(values, when=None):
    entry = {"date": when or date.today().isoformat()}
    for f in STAT_FIELDS:
        if values.get(f) not in (None, ""):
            entry[f] = _num(values[f])
    if len(entry) == 1:
        return {"error": "Type at least one number."}
    stats = [s for s in get_stats() if s.get("date") != entry["date"]] + [entry]
    _write(STATS_FILE, stats)
    return entry


# ------------------------------------------------------------------ calories (BMI / TDEE)
PROFILE_FILE = data_file("gym_profile.json")
KCAL_PER_KG = 7700          # about how many calories one kilo of body weight is
ACTIVITY = [                 # how hard you train in a week -> how much to multiply your resting calories by
    ("none", "No training (desk job)", 1.2),
    ("light", "Light: 1-3 workouts a week", 1.375),
    ("moderate", "Moderate: 3-5 workouts a week", 1.55),
    ("hard", "Hard: 6-7 workouts a week", 1.725),
    ("athlete", "Very hard: twice a day / physical job", 1.9),
]


def bmi_category(bmi):
    return ("Underweight" if bmi < 18.5 else "Healthy" if bmi < 25 else "Overweight" if bmi < 30 else "Obese")


def calories(height_cm, weight_kg, age, gender, activity="moderate"):
    """BMI, resting calories (Mifflin-St Jeor) and the calories to cut, keep or bulk,
    with how fast your weight should change on each."""
    h, w, a = _num(height_cm), _num(weight_kg), _num(age)
    if not h or not w or not a or h < 80 or w < 25 or not 10 <= a <= 100:
        return {"error": "Type your height (cm), weight (kg) and age."}
    male = str(gender or "").lower().startswith("m")
    bmr = 10 * w + 6.25 * h - 5 * a + (5 if male else -161)
    mult = next((m for k, _, m in ACTIVITY if k == activity), 1.55)
    keep = bmr * mult
    bmi = w / (h / 100) ** 2

    def plan(key, label, delta):
        kcal = round((keep + delta) / 10) * 10
        week = delta * 7 / KCAL_PER_KG
        return {"key": key, "label": label, "kcal": kcal, "kg_week": round(week, 2), "kg_month": round(week * 30 / 7, 1),
                "weight_in_3_months": round(w + week * 13, 1)}

    return {
        "bmi": round(bmi, 1), "category": bmi_category(bmi), "bmr": round(bmr), "maintenance": round(keep / 10) * 10,
        "healthy_weight": [round(18.5 * (h / 100) ** 2, 1), round(24.9 * (h / 100) ** 2, 1)],
        "protein_g": [round(w * 1.6), round(w * 2.2)],
        "plans": [plan("cut", "Deficit (lose fat)", -500), plan("mild_cut", "Small deficit", -250),
                  plan("keep", "Maintenance", 0), plan("lean_bulk", "Lean bulk", 250), plan("bulk", "Bulk (gain)", 500)],
    }


def get_profile():
    p = _read(PROFILE_FILE, {})
    return p if isinstance(p, dict) else {}


def save_profile(profile):
    keep = {k: profile.get(k) for k in ("height", "weight", "age", "gender", "activity") if profile.get(k) not in (None, "")}
    _write(PROFILE_FILE, keep)
    result = calories(keep.get("height"), keep.get("weight"), keep.get("age"), keep.get("gender"), keep.get("activity", "moderate"))
    if not result.get("error") and keep.get("weight"):
        stats = get_stats()                        # the weight also goes into Body stats (once a day)
        today = date.today().isoformat()
        if not any(s.get("date") == today and s.get("weight") == _num(keep["weight"]) for s in stats):
            add_stats({**next((s for s in stats if s.get("date") == today), {}), "weight": keep["weight"]})
    return {"profile": keep, **result}


# ------------------------------------------------------------------ meal planner (bulk / cut / recomp)
# per 100 g: kcal, protein, fat, carbs  (cooked weights for meat, rice and potato)
FOODS = {
    "Eggs": (143, 12.6, 9.5, 0.7), "Oats": (389, 17, 7, 66), "Banana": (89, 1.1, 0.3, 23),
    "Chicken breast": (165, 31, 3.6, 0), "Rice": (130, 2.7, 0.3, 28), "Vegetables": (34, 2.8, 0.4, 7),
    "Olive oil": (884, 0, 100, 0), "Greek yogurt": (59, 10, 0.4, 3.6), "Almonds": (579, 21, 50, 22),
    "Whey protein": (400, 80, 6, 8), "Milk": (61, 3.2, 3.3, 4.8), "Peanut butter": (588, 25, 50, 20),
    "Salmon": (208, 20, 13, 0), "Lean beef": (250, 26, 15, 0), "Tuna (in water)": (116, 26, 1, 0),
    "Sweet potato": (86, 1.6, 0.1, 20), "Whole-wheat bread": (247, 13, 3.4, 41), "Tofu": (144, 17, 9, 3),
    "Pasta": (158, 5.8, 0.9, 31),
}
GOALS = {   # kcal vs maintenance, protein g/kg, fat g/kg
    "bulk": {"label": "Bulk", "delta": 300, "protein": 1.8, "fat": 1.0,
             "note": "A small surplus: you gain muscle with as little fat as possible (about +0.25 kg a week)."},
    "cut": {"label": "Cut", "delta": -500, "protein": 2.2, "fat": 0.8,
            "note": "A deficit with high protein so you keep your muscle while losing fat (about -0.5 kg a week)."},
    "recomp": {"label": "Recomposition", "delta": 0, "protein": 2.2, "fat": 0.9,
               "note": "Eat at maintenance with high protein and train hard: slowly lose fat and gain muscle at the same weight."},
}
# Dishes: (name, protein food, carb food, fat food or None, fixed extras {food: grams}, how to make it).
# The grams of the first three are worked out so each meal hits its share of the day's protein, carbs and fats.
DISHES = {
    "breakfast": [
        ("Scrambled eggs on toast", "Eggs", "Whole-wheat bread", "Olive oil", {"Vegetables": 80},
         "Scramble the eggs in the oil, add spinach or tomato, serve on toasted bread."),
        ("Overnight oats", "Greek yogurt", "Oats", "Peanut butter", {"Banana": 100, "Milk": 150},
         "Stir oats, milk and yogurt in a jar the night before; top with banana and peanut butter."),
        ("Protein oatmeal", "Whey protein", "Oats", "Almonds", {"Milk": 200, "Banana": 80},
         "Cook the oats in the milk, stir in the protein powder off the heat, top with banana and almonds."),
        ("Veggie omelette & sweet potato", "Eggs", "Sweet potato", "Olive oil", {"Vegetables": 120},
         "Pan-fry diced sweet potato, then pour the beaten eggs over the vegetables and fold."),
    ],
    "lunch": [
        ("Chicken rice bowl", "Chicken breast", "Rice", "Olive oil", {"Vegetables": 150},
         "Grill or pan-fry the chicken, slice it over rice with vegetables and a drizzle of oil."),
        ("Tuna sandwich & salad", "Tuna (in water)", "Whole-wheat bread", "Olive oil", {"Vegetables": 120},
         "Mix tuna with a little oil, lemon and pepper; make a sandwich and have a side salad."),
        ("Beef & sweet potato bowl", "Lean beef", "Sweet potato", "Olive oil", {"Vegetables": 150},
         "Brown the beef with garlic, serve with roasted sweet potato and greens."),
        ("Tofu stir-fry", "Tofu", "Rice", "Olive oil", {"Vegetables": 200},
         "Fry tofu cubes until golden, add vegetables and soy sauce, serve over rice."),
    ],
    "dinner": [
        ("Salmon & sweet potato", "Salmon", "Sweet potato", None, {"Vegetables": 150},
         "Bake the salmon and sweet potato at 200 C for about 20 minutes, with steamed vegetables."),
        ("Chicken adobo with rice", "Chicken breast", "Rice", "Olive oil", {"Vegetables": 100},
         "Simmer chicken in soy sauce, vinegar, garlic and bay leaf until tender; serve with rice."),
        ("Beef stir-fry", "Lean beef", "Rice", "Olive oil", {"Vegetables": 200},
         "Stir-fry thin beef strips on high heat, add vegetables and sauce, serve over rice."),
        ("Chicken pasta", "Chicken breast", "Pasta", "Olive oil", {"Vegetables": 120},
         "Toss cooked pasta with pan-fried chicken, vegetables, garlic and a little oil."),
    ],
    "snack": [
        ("Greek yogurt & banana", "Greek yogurt", "Banana", "Almonds", {},
         "Yogurt in a bowl, sliced banana and a few almonds on top."),
        ("Protein shake", "Whey protein", "Banana", "Peanut butter", {"Milk": 250},
         "Blend milk, protein powder, banana and peanut butter."),
        ("Tuna on toast", "Tuna (in water)", "Whole-wheat bread", None, {},
         "Tuna with pepper and lemon on toasted bread."),
        ("Boiled eggs & fruit", "Eggs", "Banana", None, {},
         "Hard-boil a few eggs (10 minutes) and have them with a banana."),
    ],
}
# how the day is split for each number of meals: (name, kind of dish, share of the day)
MEAL_SPLITS = {
    2: [("Lunch", "lunch", 0.5), ("Dinner", "dinner", 0.5)],
    3: [("Breakfast", "breakfast", 0.3), ("Lunch", "lunch", 0.4), ("Dinner", "dinner", 0.3)],
    4: [("Breakfast", "breakfast", 0.25), ("Lunch", "lunch", 0.35), ("Snack", "snack", 0.1), ("Dinner", "dinner", 0.3)],
    5: [("Breakfast", "breakfast", 0.22), ("Morning snack", "snack", 0.1), ("Lunch", "lunch", 0.3),
        ("Afternoon snack", "snack", 0.1), ("Dinner", "dinner", 0.28)],
    6: [("Breakfast", "breakfast", 0.2), ("Morning snack", "snack", 0.1), ("Lunch", "lunch", 0.25),
        ("Afternoon snack", "snack", 0.1), ("Dinner", "dinner", 0.25), ("Evening snack", "snack", 0.1)],
}


def _macros(items):
    t = [0.0, 0.0, 0.0, 0.0]
    for food, g in items.items():
        for i, v in enumerate(FOODS[food]):
            t[i] += v * g / 100
    return {"kcal": round(t[0]), "protein": round(t[1]), "fat": round(t[2]), "carbs": round(t[3])}


def meal_plan(goal="recomp", profile=None, meals=4, dishes=None):
    """A day of meals for bulk / cut / recomp, from the Calories tab numbers: calories, protein, fat and carbs
    for the day, split over 2-6 meals. Each meal is a dish with gram amounts (estimates).
    dishes: optional list with the dish number for each meal (the "Other dish" button changes one)."""
    goal = goal if goal in GOALS else "recomp"
    try:
        meals = min(6, max(2, int(meals)))
    except (TypeError, ValueError):
        meals = 4
    p = profile if profile is not None else get_profile()
    base = calories(p.get("height"), p.get("weight"), p.get("age"), p.get("gender"), p.get("activity", "moderate"))
    if base.get("error"):
        return {"error": "Fill in the Calories tab first (height, weight, age) - the meal plan is made from it."}
    g, w = GOALS[goal], _num(p["weight"])
    kcal = round((base["maintenance"] + g["delta"]) / 10) * 10
    protein, fat = round(w * g["protein"]), round(w * g["fat"])
    carbs = max(50, round((kcal - protein * 4 - fat * 9) / 4))
    picks = list(dishes or [])
    out, seen = [], {}
    for i, (name, kind, share) in enumerate(MEAL_SPLITS[meals]):
        options = DISHES[kind]
        n = seen.get(kind, 0)                              # two snacks in a day get different dishes
        seen[kind] = n + 1
        try:
            pick = int(picks[i]) % len(options)
        except (IndexError, TypeError, ValueError):
            pick = n % len(options)
        dish, pf, cf, ff, extra, how = options[pick]
        items = dict(extra)
        want = {"protein": protein * share, "fat": fat * share, "carbs": carbs * share}
        now = _macros(items)
        items[cf] = items.get(cf, 0) + max(0, (want["carbs"] - now["carbs"]) / FOODS[cf][3] * 100)     # carbs first
        now = _macros(items)
        items[pf] = items.get(pf, 0) + max(0, (want["protein"] - now["protein"]) / FOODS[pf][1] * 100)  # then protein
        if ff:                                                                                         # fat tops up the rest
            now = _macros(items)
            items[ff] = max(0, (want["fat"] - now["fat"]) / FOODS[ff][2] * 100)
        items = {f: int(round(v / 5) * 5) for f, v in items.items() if v >= 5}
        out.append({"name": name, "dish": dish, "how": how, "pick": pick, "choices": len(options),
                    "items": [{"food": f, "grams": v} for f, v in items.items()], **_macros(items)})
    totals = {k: sum(m[k] for m in out) for k in ("kcal", "protein", "fat", "carbs")}
    return {"goal": goal, "label": g["label"], "note": g["note"], "kcal": kcal, "protein": protein, "fat": fat,
            "carbs": carbs, "maintenance": base["maintenance"], "meals": out, "meal_count": meals, "totals": totals,
            "goals": [[k, v["label"]] for k, v in GOALS.items()], "meal_options": sorted(MEAL_SPLITS)}


# ------------------------------------------------------------------ monthly lifting graphs
def e1rm(kg, reps):
    """estimated one-rep max (Epley): 60 kg x 10 -> 80 kg"""
    kg, reps = _num(kg) or 0, rep_value(reps) or 0
    if not kg or not reps or reps > 30:
        return kg or 0
    return round(kg * (1 + reps / 30), 1)


def monthly_progress():
    """Per exercise, per month: heaviest set (weight + the reps done with it), best estimated 1RM, total volume,
    number of sets. Plus every Body stats entry (weight, arms, thighs, chest, waist). Oldest first."""
    months = {}
    for d, entry in _log().items():
        m = d[:7]
        for st in entry.get("sets", []):
            kg, reps = _num(st.get("kg")) or 0, rep_value(st.get("reps")) or 0
            row = months.setdefault(st["exercise"], {}).setdefault(m, {"month": m, "best": 0, "reps": 0, "e1rm": 0, "volume": 0, "sets": 0, "days": set()})
            if kg > row["best"] or (kg == row["best"] and reps > row["reps"]):
                row["reps"] = reps
            row["best"] = max(row["best"], kg)
            row["e1rm"] = max(row["e1rm"], e1rm(kg, st.get("reps")))
            row["volume"] += kg * reps
            row["sets"] += 1
            row["days"].add(d)
    out = {}
    for ex, by in months.items():
        rows = [{**r, "volume": round(r["volume"]), "days": len(r["days"])} for _, r in sorted(by.items())]
        out[ex] = rows
    weights = {}
    for s in get_stats():
        if s.get("weight"):
            weights.setdefault(s["date"][:7], []).append(s["weight"])
    body = [{"month": m, "weight": round(sum(v) / len(v), 1)} for m, v in sorted(weights.items()) if re.match(r"\d{4}-\d\d", m)]
    stats = [{"label": s.get("date", ""), **{f: s.get(f) for f in STAT_FIELDS}} for s in get_stats()]
    return {"exercises": dict(sorted(out.items(), key=lambda kv: -sum(r["sets"] for r in kv[1]))), "body": body,
            "stats": stats, "stat_fields": STAT_FIELDS}


# ------------------------------------------------------------------ text for the AI
def plan_text(day_num=None):
    p = day_plan(day_num)
    if not p:
        return "No gym routine yet. Import your Excel routine in the Gym panel."
    lines = [f"Day {p['day']}: {p['title']}"]
    for ex in p["exercises"]:
        w = f" @ {ex['kg']:g} kg" if ex.get("kg") else ""
        lines.append(f"- {ex['name']}: {ex['sets']} x {ex['target']}{w}. {ex['tip']}".strip())
    return "\n".join(lines)


def log_text(exercise, set_no, kg=None, reps=None):
    e = log_set(exercise, set_no, kg, reps, day=next_day())
    return f"Logged {exercise} set {set_no}: {kg or '-'} kg x {reps or '-'}. {len(e['sets'])} sets done today."


def history_text():
    h = history(14)
    if not h:
        return "No workouts logged in the last 2 weeks."
    return f"Streak: {streak()} days.\n" + "\n".join(
        f"{w['date']}: {w['title'] or 'Day ' + str(w['day'])} - {w['exercises']} exercises, {w['sets']} sets, volume {w['volume']} kg"
        for w in h)


# ---------------- exercise alternatives (Gym -> Alternatives tab) ----------------
EXERCISE_DB = os.path.join(os.path.dirname(os.path.abspath(__file__)), "exercise_db.json")
_db_cache = None


def exercise_db():
    global _db_cache
    if _db_cache is None:
        with open(EXERCISE_DB, encoding="utf-8") as f:
            _db_cache = json.load(f)
    return _db_cache


def _norm(s):
    return re.sub(r"[^a-z0-9 ]+", " ", (s or "").lower().replace("-", " ")).split()


def _match_score(query, ex):
    """how well a search matches an exercise: 0-100"""
    from rapidfuzz import fuzz
    q = " ".join(_norm(query))
    best = 0
    for name in [ex["name"]] + ex.get("aliases", []):
        n = " ".join(_norm(name))
        best = max(best, fuzz.ratio(q, n), fuzz.token_set_ratio(q, n) - 4, 100 if q == n else 0)
    return best


def _public(ex, group, source="built-in"):
    return {"name": ex["name"], "stars": max(0, min(5, int(ex.get("stars", 0)))), "used_by": ex.get("used_by", []),
            "equipment": ex.get("equipment", ""), "why": ex.get("why", ""), "group": group["name"],
            "muscles": group.get("muscles", []), "source": source}


# which muscles an exercise works, most-worked first (for the muscle map next to each exercise)
MUSCLE_RULES = [
    (r"incline.*(press|bench|fly)", ["upper chest", "front delts", "triceps"]),
    (r"decline.*(press|bench)", ["lower chest", "triceps", "front delts"]),
    (r"bench|chest press|push ?up|dip|pec|fly|flye|crossover", ["chest", "triceps", "front delts"]),
    (r"overhead press|shoulder press|military|arnold|ohp", ["front delts", "side delts", "triceps", "upper back"]),
    (r"lateral raise|side raise|upright row", ["side delts", "traps", "front delts"]),
    (r"rear delt|reverse fly|face ?pull", ["rear delts", "upper back", "traps"]),
    (r"pull ?up|chin ?up|pulldown|pull down|lat ", ["lats", "biceps", "upper back", "rear delts"]),
    (r"row", ["upper back", "lats", "rear delts", "biceps"]),
    (r"shrug", ["traps", "forearms"]),
    (r"deadlift|rdl|good morning|back extension|hyperextension", ["hamstrings", "glutes", "lower back", "upper back"]),
    (r"hip thrust|glute bridge|kickback|abduction", ["glutes", "hamstrings", "quads"]),
    (r"squat|leg press|lunge|split squat|step ?up|hack", ["quads", "glutes", "adductors", "hamstrings"]),
    (r"leg extension", ["quads"]),
    (r"leg curl|hamstring curl|nordic", ["hamstrings", "calves"]),
    (r"calf|calves", ["calves"]),
    (r"curl", ["biceps", "forearms"]),
    (r"tricep|pushdown|skull ?crusher|close.?grip|kickback|extension", ["triceps", "front delts"]),
    (r"crunch|sit ?up|plank|leg raise|ab |abs|hollow|rollout", ["abs", "obliques"]),
    (r"twist|woodchop|side bend", ["obliques", "abs"]),
    (r"farmer|wrist|grip", ["forearms", "traps"]),
]


def muscles_for(name):
    """'Incline Bench' -> ['upper chest', 'front delts', 'triceps'] (most targeted first, at most 4)."""
    n = " " + " ".join(_norm(name)) + " "
    for pat, ms in MUSCLE_RULES:                        # specific patterns first (incline bench = UPPER chest)
        if re.search(pat, n):
            return ms[:4]
    best, best_g = 0, None
    for g in exercise_db()["groups"]:
        for ex in g["exercises"]:
            try:
                s = _match_score(name, ex)
            except Exception:
                s = 0
            if s > best: best, best_g = s, g
    return (best_g or {}).get("muscles", [])[:4] if best >= 70 else []


def find_alternatives(query, use_ai=True):
    """'bench press' -> that exercise + every other exercise that trains the same thing, best (5 stars) first.
    A muscle ('chest', 'lats') -> every exercise for it. Unknown exercises -> the AI suggests some (marked as AI)."""
    db = exercise_db()
    query = (query or "").strip()
    if not query:
        return {"error": "Type an exercise (e.g. 'bench press') or a muscle (e.g. 'chest')."}
    qn = " ".join(_norm(query))
    # a muscle name: every group that trains it
    groups = [g for g in db["groups"] if any(qn == m or qn.rstrip("s") == m.rstrip("s") for m in g.get("muscles", []))]
    exact = any(qn == " ".join(_norm(n)) for g in db["groups"] for ex in g["exercises"] for n in [ex["name"]] + ex.get("aliases", []))
    if groups and not exact:
        out = [_public(ex, g) for g in groups for ex in g["exercises"]]
        out.sort(key=lambda e: (-e["stars"], e["name"]))
        return {"query": query, "kind": "muscle", "title": f"Exercises for {query}", "match": None, "alternatives": out,
                "people": db["people"]}
    scored = sorted(((_match_score(query, ex), g, ex) for g in db["groups"] for ex in g["exercises"]), key=lambda x: -x[0])
    if scored and scored[0][0] >= 72:
        _, g, ex = scored[0]
        alts = [_public(x, g) for x in g["exercises"] if x is not ex]
        alts.sort(key=lambda e: (-e["stars"], e["name"]))
        return {"query": query, "kind": "exercise", "title": f"Alternatives to {ex['name']}", "match": _public(ex, g),
                "alternatives": alts, "people": db["people"]}
    if use_ai:
        ai = _ai_alternatives(query)
        if ai:
            return ai
    near = [ex["name"] for s, _, ex in scored[:5] if s >= 45]
    return {"query": query, "kind": "none", "title": f"No match for '{query}'", "match": None, "alternatives": [],
            "people": db["people"], "suggest": near}


def _ai_alternatives(query):
    try:
        import llm
        people = ", ".join(exercise_db()["people"])
        d = llm.ask(f"The user wants alternatives for the gym exercise '{query}'. Give 6 alternatives that train the same muscles, "
                    f"from best to worst. For each: stars 0-5 (5 = most used by well-known coaches and the best swap, 0 = weakest), "
                    f"used_by = only people from this list who are publicly known to use or recommend it ({people}) - leave it empty "
                    f"if unsure, never guess. JSON: {{\"exercise\": \"<proper name>\", \"muscles\": [\"...\"], \"alternatives\": "
                    f"[{{\"name\": \"...\", \"stars\": 4, \"used_by\": [], \"equipment\": \"...\", \"why\": \"<max 12 words>\"}}]}}",
                    model="chat", json_mode=True)
        group = {"name": d.get("exercise") or query, "muscles": [str(m) for m in d.get("muscles", [])][:4]}
        allowed = set(exercise_db()["people"])
        alts = []
        for a in d.get("alternatives", [])[:8]:
            if not isinstance(a, dict) or not a.get("name"):
                continue
            a["used_by"] = [p for p in a.get("used_by", []) if p in allowed]
            alts.append(_public(a, group, source="ai"))
        alts.sort(key=lambda e: -e["stars"])
        if alts:
            return {"query": query, "kind": "ai", "title": f"Alternatives to {group['name']} (AI suggestions - double-check)",
                    "match": None, "alternatives": alts, "people": exercise_db()["people"]}
    except Exception as e:
        print(f"[gym alternatives] {e}")
    return None


def alternatives_text(query):
    """for voice: 'what can I do instead of bench press?'"""
    r = find_alternatives(query)
    if not r.get("alternatives"):
        return r.get("error") or f"I don't know '{query}' yet." + (f" Did you mean {', '.join(r.get('suggest', []))}?" if r.get("suggest") else "")
    top = r["alternatives"][:4]
    return r["title"] + ": " + "; ".join(f"{a['name']} ({a['stars']}★{', used by ' + ', '.join(a['used_by'][:2]) if a['used_by'] else ''})" for a in top) + "."


if __name__ == "__main__":
    if len(sys.argv) > 1:
        print(import_xlsx(sys.argv[1]))
    else:
        print(__doc__)
