"""Batch 3: internet & productivity (network calls are simulated)."""
import os, sys, json, shutil, tempfile, unittest
from datetime import date, datetime, timedelta
from types import SimpleNamespace
from unittest import mock
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import folders  # noqa: E402,F401  (finds brain/, senses/, skills/, scripts/)
import sandbox  # noqa: E402,F401  (keeps the tests away from your real Spotify login, keys and installs)
import web, productivity as prod, llm, f1  # noqa: E402


class R:
    def __init__(self, text="", content=None, status=200, data=None, headers=None):
        self.text, self.status_code, self._d = text, status, data
        self.content = content if content is not None else text.encode()
        self.headers = headers or {"content-type": "text/html"}
        self.encoding = "utf-8"

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(self.status_code)

    def json(self):
        return self._d


DDG = '''<div><a rel="nofollow" class="result__a" href="//duckduckgo.com/l/?uddg=https%3A%2F%2Fexample.com%2Fcuring&rut=x">Concrete <b>curing</b></a>
<a class="result__snippet" href="#">Curing takes <b>28 days</b>.</a></div>
<div><a class="result__a" href="https://duckduckgo.com/y.js?ad=1">Ad</a></div>
<div><a class="result__a" href="https://site2.org/x">Second</a><a class="result__snippet">More.</a></div>'''
PAGE = "<html><head><title>Curing guide</title><script>var x=1</script></head><body><nav>menu</nav><p>Concrete gains " + \
       "strength over 28 days. " * 30 + "</p></body></html>"
RSS = """<rss><channel><item><title>Big news - Rappler</title><link>https://n/1</link><source>Rappler</source>
<pubDate>Tue, 20 Oct 2026</pubDate></item><item><title>Other - GMA</title><link>https://n/2</link><source>GMA</source></item></channel></rss>"""


class WebTests(unittest.TestCase):
    def test_html_to_text_skips_scripts_and_menus(self):
        title, text = web.html_to_text(PAGE)
        self.assertEqual(title, "Curing guide")
        self.assertNotIn("var x", text)
        self.assertNotIn("menu", text)

    def test_search_parses_and_skips_ads(self):
        with mock.patch.object(web.requests, "post", return_value=R(DDG)), \
                mock.patch.object(web, "ENGINES", [("duckduckgo", web._ddg_html)]):          # no real internet in tests
            res = web.search("concrete curing")
        self.assertEqual([r["url"] for r in res], ["https://example.com/curing", "https://site2.org/x"])
        self.assertEqual(res[0]["title"], "Concrete curing")
        self.assertEqual(res[0]["snippet"], "Curing takes 28 days.")

    def test_web_answer_reads_pages_and_cites(self):
        with mock.patch.object(web.requests, "post", return_value=R(DDG)), \
                mock.patch.object(web, "ENGINES", [("duckduckgo", web._ddg_html)]), \
                mock.patch.object(web.requests, "get", return_value=R(PAGE)), \
                mock.patch.object(llm, "ask", return_value="It takes 28 days [1].") as ask:
            out = web.web_answer_text("how long does concrete cure")
        self.assertIn("[1] Curing guide - https://example.com/curing", out)
        self.assertIn("ONLY these sources", ask.call_args.args[0])

    def test_news_from_rss(self):
        with mock.patch.object(web.requests, "get", return_value=R(content=RSS.encode())):
            items = web.news("Philippines")
        self.assertEqual(items[0]["title"], "Big news")
        self.assertEqual(items[0]["source"], "Rappler")

    def test_prices(self):
        self.assertEqual(web.find_price('<script>{"price": "1,299.00"}</script>'), 1299.0)
        self.assertEqual(web.find_price('<meta itemprop="price" content="499.5">'), 499.5)
        self.assertEqual(web.find_price("<p>Now only ₱2,450.00!</p>"), 2450.0)
        tmp = tempfile.mkdtemp()
        with mock.patch.object(web, "PRICES_FILE", os.path.join(tmp, "p.json")):
            with mock.patch.object(web, "check_price", return_value=("Whey 2kg", 3000.0)):
                self.assertIn("Tracking Whey 2kg", web.track_price("https://shop/whey", 2800))
            items = json.load(open(web.PRICES_FILE))
            items[0]["history"][0]["date"] = "2026-01-01"
            json.dump(items, open(web.PRICES_FILE, "w"))
            with mock.patch.object(web, "check_price", return_value=("Whey 2kg", 2750.0)):
                alerts = web.check_prices()
            self.assertIn("target", alerts[0])
        shutil.rmtree(tmp)

    def test_forecast_and_alerts(self):
        data = {"location": {"name": "Manila", "country": "Philippines"}, "current": {"temp_c": 31, "condition": {"text": "Sunny"}},
                "forecast": {"forecastday": [{"date": "2026-10-21", "day": {"condition": {"text": "Rain"}, "mintemp_c": 25,
                                                                             "maxtemp_c": 30, "daily_chance_of_rain": 90}}]},
                "alerts": {"alert": [{"headline": "Typhoon signal no. 2", "severity": "Severe"}]}}
        with mock.patch.dict(os.environ, {"WEATHER_API_KEY": "k"}), \
                mock.patch.object(web.requests, "get", return_value=R(data=data)):
            text = web.forecast_text("Manila")
        self.assertIn("90% chance of rain", text)
        self.assertIn("ALERT: Typhoon signal no. 2", text)

    def test_youtube_captions(self):
        info = {"title": "Bench tips", "channel": "Coach", "automatic_captions": {"en": [{"ext": "vtt", "url": "v"}, {"ext": "json3", "url": "j"}]}}
        self.assertEqual(web._captions(info), "j")
        self.assertIsNone(web._captions({"subtitles": {}}))


ICS = """BEGIN:VCALENDAR
BEGIN:VEVENT
SUMMARY:Thesis meeting\\, room 3
DTSTART:{d}T060000Z
DTEND:{d}T070000Z
END:VEVENT
BEGIN:VEVENT
SUMMARY:Gym
DTSTART;TZID=Asia/Manila:{d}T180000
RRULE:FREQ=WEEKLY;COUNT=3
END:VEVENT
BEGIN:VEVENT
SUMMARY:Holiday
DTSTART;VALUE=DATE:{d}
END:VEVENT
END:VCALENDAR"""


class ProductivityTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.saved = json.loads(json.dumps(f1.user_data))
        f1.user_data.update({"events": [], "tasks": [], "schedule": {}, "reminders": []})
        self.p = [mock.patch.object(prod, "STATE_FILE", os.path.join(self.tmp, "s.json")),
                  mock.patch.object(prod, "DOCS_DIR", os.path.join(self.tmp, "docs")),
                  mock.patch.object(prod, "MEETINGS_DIR", os.path.join(self.tmp, "meet")),
                  mock.patch.object(f1, "save_data", lambda d: None)]
        for p in self.p: p.start()

    def tearDown(self):
        for p in self.p: p.stop()
        f1.user_data.clear(); f1.user_data.update(self.saved)
        shutil.rmtree(self.tmp)

    def test_ics_parse_and_sync(self):
        d = (date.today() + timedelta(days=1)).strftime("%Y%m%d")
        events = prod.parse_ics(ICS.format(d=d))
        titles = [e["title"] for e in events]
        self.assertEqual(titles.count("Gym"), 3)
        self.assertIn("Thesis meeting, room 3", titles)
        hol = next(e for e in events if e["title"] == "Holiday")
        self.assertIsNone(hol["time"])
        with mock.patch.object(prod.requests, "get", return_value=R(ICS.format(d=d))):
            self.assertIn("Synced 5", prod.sync_ical("https://calendar.google.com/x/basic.ics"))
            prod.sync_ical()                                  # second sync replaces, doesn't duplicate
        self.assertEqual(len([e for e in f1.user_data["events"] if e.get("source") == "google"]), 5)
        self.assertIn("Import", prod.export_ical())

    def test_free_slots_and_scheduling(self):
        day = (date.today() + timedelta(days=1)).isoformat()
        f1.user_data["events"] = [{"title": "Class", "date": day, "time": "08:00", "end_time": "12:00"},
                                  {"title": "Lunch", "date": day, "time": "12:00", "end_time": "13:00"}]
        slots = prod.free_slots(day, 90)
        self.assertEqual(slots[0]["start"], "13:00")
        msg = prod.schedule_task("Thesis writing", 120, day)
        self.assertIn("at 13:00", msg)
        self.assertEqual(prod.free_slots(day, 90)[0]["start"], "15:00")

    def test_briefing_once_a_day(self):
        self.assertTrue(prod.briefing_due(datetime(2026, 10, 21, 7)))
        self.assertFalse(prod.briefing_due(datetime(2026, 10, 21, 9)))
        self.assertFalse(prod.briefing_due(datetime(2026, 10, 22, 4)))
        with mock.patch("web.news", return_value=[{"title": "Headline 1"}]):
            self.assertIn("Headline 1", prod.morning_briefing())

    def test_create_document(self):
        out = prod.create_document("Weekly Report", "# Summary\nAll good.\n- **Bench** 22.5 kg\n1. Rest", "docx", open_it=False)
        path = out.split("Created ")[1].rstrip(".")
        from docx import Document
        doc = Document(path)
        self.assertEqual(doc.paragraphs[0].text, "Weekly Report")
        self.assertTrue(any(r.bold and r.text == "Bench" for p in doc.paragraphs for r in p.runs))
        self.assertIn("(2)", prod.create_document("Weekly Report", "x", "txt", open_it=False) +
                      prod.create_document("Weekly Report", "x", "txt", open_it=False))

    def test_inbox_summary_read_only(self):
        raw = b"From: Prof <p@x.edu>\r\nSubject: Thesis due\r\n\r\nPlease submit chapter 3 by Friday."
        box = mock.MagicMock()
        box.search.return_value = ("OK", [b"1"])
        box.fetch.return_value = ("OK", [(b"1", raw)])
        with mock.patch.dict(os.environ, {"EMAIL_ADDRESS": "a@gmail.com", "EMAIL_APP_PASSWORD": "pw"}), \
                mock.patch.object(prod.imaplib, "IMAP4_SSL", return_value=box), \
                mock.patch.object(llm, "ask", return_value="Needs action: Thesis due Friday") as ask:
            self.assertIn("Thesis due", prod.inbox_summary())
        box.select.assert_called_with("INBOX", readonly=True)
        self.assertIn("chapter 3 by Friday", ask.call_args.args[0])

    def test_meeting_notes(self):
        import numpy as np, stt
        prod.recorder.stream = mock.Mock()
        prod.recorder.frames = [np.zeros((16000, 1), dtype="int16")]
        prod.recorder.started, prod.recorder.title = datetime.now().timestamp() - 120, "Thesis group"
        notes = {"summary": "We split the survey.", "decisions": ["Use Google Forms"],
                 "action_items": [{"task": "Send survey", "who": "me", "due": ""}, {"task": "Check data", "who": "Ana"}]}
        with mock.patch.object(stt, "transcribe", return_value="okay let's split the survey"), \
                mock.patch.object(llm, "ask", return_value=notes), mock.patch.object(f1, "add_task") as add:
            r = prod.stop_meeting()
        self.assertEqual((r["minutes"], r["tasks_added"]), (2, 1))
        add.assert_called_once()
        self.assertIn("Use Google Forms", open(r["notes_file"], encoding="utf-8").read())
        self.assertIn("error", prod.stop_meeting())


if __name__ == "__main__":
    unittest.main()
