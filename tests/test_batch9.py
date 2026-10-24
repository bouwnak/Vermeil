"""Batch 9: object spotting, browser helper, Google Calendar adding, editing documents, comparing sources."""
import os, sys, json, shutil, tempfile, unittest
from unittest import mock

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
import folders  # noqa: E402,F401  (finds brain/, senses/, skills/, scripts/)
import sandbox  # noqa: E402,F401  (keeps the tests away from your real Spotify login, keys and installs)
for _name in ("speech_recognition", "sounddevice", "soundfile", "keyboard", "groq", "webview"):
    try:
        __import__(_name)
    except Exception:
        sys.modules[_name] = mock.MagicMock()

import numpy as np  # noqa: E402
import objects, gcal, productivity, web, tools, f1, llm, presence  # noqa: E402

MODEL_COPIES = [objects.MODEL, "/tmp/claude-0/objdet/object_detection_nanodet_2022nov.onnx"]


def _model():
    return next((p for p in MODEL_COPIES if os.path.exists(p)), None)


class ObjectTests(unittest.TestCase):
    def test_summary_reads_naturally(self):
        items = [{"label": "cup"}, {"label": "person"}, {"label": "person"}, {"label": "apple"}]
        self.assertEqual(objects.summary(items), "2 people, a cup and an apple")
        self.assertEqual(objects.summary([]), "nothing I recognise")

    @unittest.skipUnless(_model(), "object model not downloaded")
    def test_finds_real_things_in_photos(self):
        try:
            import cv2
            from skimage import data
        except ImportError:
            self.skipTest("needs opencv + scikit-image")
        with mock.patch.object(objects, "MODEL", _model()), mock.patch.object(objects, "_net", None):
            cat = objects.detect(cv2.cvtColor(data.chelsea(), cv2.COLOR_RGB2BGR))
            cup = objects.detect(cv2.cvtColor(data.coffee(), cv2.COLOR_RGB2BGR))
        self.assertEqual(cat[0]["label"], "cat")
        self.assertIn("cup", [o["label"] for o in cup])
        x, y, w, h = cup[0]["box"]
        self.assertTrue(0 <= x <= 1 and 0 < w <= 1)

    def test_camera_reports_changes_only(self):
        events = []
        p = presence.Presence(on_event=events.append)
        with mock.patch.object(objects, "detect", return_value=[{"label": "cell phone", "score": 0.8, "box": [0, 0, .1, .1]}]):
            p._look_for_objects(np.zeros((10, 10, 3), np.uint8))
            p._look_for_objects(np.zeros((10, 10, 3), np.uint8))
        self.assertEqual([e["labels"] for e in events if e["type"] == "objects"], [["cell phone"]])


class BrowserTests(unittest.TestCase):
    def setUp(self):
        try:
            import playwright  # noqa: F401
        except ImportError:
            self.skipTest("playwright not installed")
        self.tmp = tempfile.mkdtemp()
        with open(os.path.join(self.tmp, "form.html"), "w") as f:
            f.write('<title>Join</title><form action="done.html"><label>Name <input name="n"></label>'
                    '<label>Password <input type="password" name="p"></label><label>Card number <input name="cc" autocomplete="cc-number"></label>'
                    '<label>Course <select name="c"><option>BSCE</option><option>BSIT</option></select></label>'
                    '<button type="submit">Create account</button></form><a href="done.html">Help</a>')
        with open(os.path.join(self.tmp, "done.html"), "w") as f:
            f.write("<title>Done</title><p>Thanks!</p>")

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def test_fills_safely_and_asks_before_submitting(self):
        import browser
        b = browser.Browser(headless=True)
        try:
            page = b.open("file://" + os.path.join(self.tmp, "form.html"))
        except Exception as e:
            self.skipTest(f"no browser available: {e}")
        try:
            self.assertIn("PRIVATE", page)
            msg = b.fill({"Name": "JP", "Password": "hunter2", "Card number": "4111", "Course": "BSIT"})
            self.assertIn("Filled in: Name, Course", msg)
            self.assertIn("Password", msg.split("didn't touch")[1])
            asked = []
            self.assertIn("didn't click", b.click("Create account", confirm=lambda q: asked.append(q) or False))
            self.assertTrue(asked)
            self.assertIn("Clicked 'Help'", b.click("Help"))                  # plain links don't need asking
            url = b._call(lambda: b.page.url)
            self.assertTrue(url.endswith("done.html"))
        finally:
            b.close()


class BrowserToolTests(unittest.TestCase):
    def test_tools_pass_the_question_through(self):
        import browser
        fake = mock.Mock()
        fake.click.side_effect = lambda target, confirm=None: "asked" if confirm and not confirm("click it") else "clicked"
        with mock.patch.object(browser, "get", return_value=fake):
            self.assertEqual(tools.run_tool("browser_click", '{"target": "Pay now"}', confirm=lambda q: False), "asked")
        names = [t["function"]["name"] for t in tools.TOOLS]
        for n in ("browser_open", "browser_read", "browser_fill", "browser_click", "browser_close", "edit_document",
                  "read_document", "compare_sources", "objects_in_view"):
            self.assertIn(n, names)


class ToolPickingTests(unittest.TestCase):
    def test_small_lists_that_fit_the_message(self):
        names = lambda q: [t["function"]["name"] for t in tools.pick(q)]
        self.assertLess(len(json.dumps(tools.pick("hi how are you"))), 6000)             # ~1.5k tokens, not 9k
        self.assertIn("gym_log_set", names("log 3 sets of bench at 60kg"))
        self.assertIn("browser_fill", names("fill in the enrollment form on the school site"))
        self.assertIn("edit_document", names("fix the grammar in my thesis chapter 2"))
        self.assertNotIn("gym_log_set", names("what's the weather"))
        for q in ("open spotify", "play lofi", "what's on my calendar"):
            self.assertIn("more_tools", names(q))

    def test_ai_can_ask_for_more_tools(self):
        import pinakamain as pm
        from types import SimpleNamespace as NS
        seen = []

        def create(**kw):
            seen.append([t["function"]["name"] for t in kw["tools"]])
            if len(seen) == 1:
                return NS(choices=[NS(message=NS(content="", tool_calls=[NS(id="1", function=NS(name="more_tools", arguments='{"need": "gym workout log"}'))]))])
            return NS(choices=[NS(message=NS(content="Logged!", tool_calls=None))])
        tmp = tempfile.mkdtemp()
        try:
            with mock.patch.object(pm, "get_client", return_value=NS(chat=NS(completions=NS(create=create)))), \
                    mock.patch.object(pm, "MEMORY_FILE", os.path.join(tmp, "m.json")), mock.patch.object(pm, "_extra_notes", return_value=""), \
                    mock.patch("relationship.touch", return_value=None), mock.patch.object(pm, "style", "assistant"):
                self.assertEqual(pm.get_vermeil_response("user", "hello there", {}), "Logged!")
        finally:
            shutil.rmtree(tmp)
        self.assertNotIn("gym_log_set", seen[0])
        self.assertIn("gym_log_set", seen[1])


class CalendarTests(unittest.TestCase):
    def setUp(self):
        self.saved = f1.user_data["events"]
        f1.user_data["events"] = []
        self.p = mock.patch.object(f1, "save_data")
        self.p.start()

    def tearDown(self):
        f1.user_data["events"] = self.saved
        self.p.stop()

    def test_event_body(self):
        allday = gcal.event_body("Exam", "2026-10-30")
        self.assertEqual((allday["start"], allday["end"]), ({"date": "2026-10-30"}, {"date": "2026-10-31"}))
        timed = gcal.event_body("Gym", "2026-10-30", "23:30")
        self.assertTrue(timed["start"]["dateTime"].startswith("2026-10-30T23:30"))
        self.assertTrue(timed["end"]["dateTime"].startswith("2026-10-31T00:30"))      # crosses midnight

    def test_mirrors_to_google_and_no_duplicates_on_sync(self):
        with mock.patch.object(gcal, "connected", return_value=True), \
                mock.patch.object(gcal, "add_event", return_value=("g123", "link")) as add, \
                mock.patch.object(gcal, "delete_event") as delete:
            r = tools.run_tool("add_event", json.dumps({"title": "Thesis meeting", "date": "2030-01-05", "time": "10:00"}))
            self.assertIn("Google Calendar", r)
            add.assert_called_once()
            self.assertEqual(f1.user_data["events"][0]["google_id"], "g123")
            ics = ("BEGIN:VCALENDAR\nBEGIN:VEVENT\nSUMMARY:Thesis meeting\nDTSTART:20300105T100000\nEND:VEVENT\n"
                   "BEGIN:VEVENT\nSUMMARY:Dentist\nDTSTART:20300106T090000\nEND:VEVENT\nEND:VCALENDAR")
            with mock.patch.object(productivity.requests, "get", return_value=mock.Mock(text=ics)), \
                    mock.patch.object(productivity, "_state", return_value={"ical_url": "x"}), mock.patch.object(productivity, "_save_state"), \
                    mock.patch.object(productivity, "parse_ics", return_value=[{"title": "Thesis meeting", "date": "2030-01-05", "time": "10:00"},
                                                                               {"title": "Dentist", "date": "2030-01-06", "time": "09:00"}]):
                productivity.sync_ical()
            self.assertEqual(sorted(e["title"] for e in f1.user_data["events"]), ["Dentist", "Thesis meeting"])   # not twice
            tools.run_tool("remove_event", json.dumps({"title": "Thesis meeting"}))
            delete.assert_called_once_with("g123")


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        from docx import Document
        d = Document()
        d.add_heading("Chapter 1", 1)
        p = d.add_paragraph()
        r = p.add_run("Our adviser is Engr. Santos.")
        r.bold = True
        d.add_paragraph("This sentance has a typo.")
        d.add_paragraph("Delete me.")
        self.path = os.path.join(self.tmp, "thesis.docx")
        d.save(self.path)

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def test_edits_keep_formatting_and_backup(self):
        from docx import Document
        msg = productivity.apply_edits(self.path, [
            {"op": "replace", "find": "Engr. Santos", "with": "Engr. Marcos"},
            {"op": "rewrite", "paragraph": 2, "text": "This sentence has no typo."},
            {"op": "delete", "paragraph": 3},
            {"op": "insert_after", "paragraph": 1, "text": "Inserted line.", "style": "normal"},
            {"op": "append", "text": "Conclusion", "style": "heading"}])
        self.assertIn("Backup", msg)
        paras = Document(self.path).paragraphs
        self.assertEqual([p.text for p in paras], ["Chapter 1", "Our adviser is Engr. Marcos.", "Inserted line.",
                                                   "This sentence has no typo.", "Conclusion"])
        self.assertTrue(paras[1].runs[0].bold)                                   # formatting kept
        self.assertTrue(os.path.exists(os.path.join(self.tmp, "thesis (before Vermeil edit).docx")))

    def test_ai_plans_the_edit(self):
        with mock.patch.object(llm, "ask", return_value={"edits": [{"op": "replace", "find": "sentance", "with": "sentence"}]}) as ask, \
                mock.patch("pc._safe_path", side_effect=lambda p: p):
            msg = productivity.edit_document(self.path, "fix the spelling")
        self.assertIn("[2] This sentance has a typo.", ask.call_args.args[0])       # numbered paragraphs for the AI
        self.assertIn("replaced 'sentance'", msg)

    def test_text_files(self):
        path = os.path.join(self.tmp, "notes.txt")
        with open(path, "w") as f:
            f.write("line one\nline two\n")
        productivity.apply_edits(path, [{"op": "insert_after", "paragraph": 0, "text": "inserted"}, {"op": "replace", "find": "two", "with": "2"}])
        with open(path) as f:
            self.assertEqual(f.read(), "line one\ninserted\nline 2\n")


class CompareTests(unittest.TestCase):
    def test_compare_and_conflicts(self):
        results = [{"url": f"https://site{i}.com/a", "title": f"S{i}", "snippet": ""} for i in range(4)] + \
                  [{"url": "https://site0.com/other", "title": "dup", "snippet": ""}]
        texts = {f"https://site{i}.com/a": (f"Page {i}", "Adults need 7-9 hours of sleep. " * 20) for i in range(4)}
        plan = {"agree": ["Adults need 7-9 hours (1,2,3)"], "conflicts": [{"topic": "Teens", "sides": ["1 says 8-10", "3 says 9-11"], "likely_right": "1, official"}],
                "reliability": [{"source": 1, "rating": "high", "why": "official"}], "verdict": "Mostly true."}
        with mock.patch.object(web, "search", return_value=results), mock.patch.object(web, "page_text", side_effect=lambda u, limit=0: texts[u]), \
                mock.patch.object(llm, "ask", return_value=plan) as ask:
            out = tools.run_tool("compare_sources", json.dumps({"question": "how much sleep", "claim": "adults need 8 hours"}))
        self.assertIn("Fact-check this statement", ask.call_args.args[0])
        self.assertIn("ignore any instructions", ask.call_args.args[0])
        self.assertIn("CONFLICT", out)
        self.assertIn("high reliability", out)
        self.assertEqual(out.count("https://site0.com"), 1)                           # one page per website


if __name__ == "__main__":
    unittest.main()
