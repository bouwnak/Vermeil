"""
Tests for study mode (academic.py), the coding helper (coder.py) and the calendar planner.
Run:  python -m unittest discover -s tests -v
The AI is replaced with a fake, so no internet or API key is needed.
"""
import os
import sys
import json
import shutil
import datetime
import tempfile
import unittest
from unittest import mock

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
import folders  # noqa: E402,F401  (finds brain/, senses/, skills/, scripts/)
import sandbox  # noqa: E402,F401  (keeps the tests away from your real Spotify login, keys and installs)
for _name in ("speech_recognition", "sounddevice", "soundfile", "keyboard", "groq", "webview"):
    try:
        __import__(_name)
    except Exception:
        sys.modules[_name] = mock.MagicMock()

import academic  # noqa: E402
import coder  # noqa: E402
import f1  # noqa: E402


def fake_reply(content=None, tool_calls=None, executed_tools=None):
    msg = mock.Mock(content=content, tool_calls=tool_calls, executed_tools=executed_tools)
    r = mock.Mock()
    r.choices = [mock.Mock(message=msg)]
    return r


def tool_call(name, args, cid="t1"):
    c = mock.Mock(id=cid)
    c.function.name = name
    c.function.arguments = json.dumps(args)
    return c


class LibraryTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        lib = os.path.join(self.tmp, "library")
        self.p = [mock.patch.object(academic, n, v) for n, v in {
            "LIB_DIR": lib, "FILES_DIR": os.path.join(lib, "files"), "TEXT_DIR": os.path.join(lib, "text"),
            "REVIEWS_DIR": os.path.join(lib, "reviews"), "INDEX_FILE": os.path.join(lib, "index.json")}.items()]
        for p in self.p:
            p.start()
        self.notes = os.path.join(self.tmp, "Beams lecture.txt")
        with open(self.notes, "w", encoding="utf-8") as f:
            f.write("Shear force is the internal force parallel to the cross-section of a beam.\n\n"
                    "The bending moment M at a section equals the sum of moments of forces on one side.\n\n"
                    "For a simply supported beam with a point load P at midspan, M_max = P*L/4.\n\n" + "Filler text. " * 300 +
                    "\n\nConcrete cover protects reinforcing steel from corrosion and fire.")

    def tearDown(self):
        for p in self.p:
            p.stop()
        shutil.rmtree(self.tmp)

    def test_add_list_search_remove(self):
        entry = academic.add_file(self.notes, "lecture", "Structural analysis")
        self.assertEqual((entry["kind"], entry["subject"]), ("lecture", "Structural analysis"))
        self.assertGreaterEqual(entry["pages"], 2)                              # long text is split into parts
        self.assertEqual(len(academic.list_files()), 1)
        hits = academic.search("maximum moment simply supported point load")
        self.assertIn("P*L/4", hits[0]["text"])
        self.assertEqual(hits[0]["label"], "Beams lecture.txt, part 1")
        self.assertIn("corrosion", academic.search("concrete cover corrosion")[0]["text"])
        self.assertEqual(academic.search("photosynthesis chlorophyll"), [])
        self.assertIn("[Beams lecture.txt, part 1]", academic.search_my_files("shear force"))
        academic.remove_file(entry["id"])
        self.assertEqual(academic.list_files(), [])
        self.assertEqual(os.listdir(academic.FILES_DIR), [])

    def test_reads_word_and_powerpoint(self):
        import docx
        from pptx import Presentation
        d = docx.Document(); d.add_paragraph("Hooke's law: stress equals E times strain."); d.save(os.path.join(self.tmp, "a.docx"))
        prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = "Torsion"
        s.placeholders[1].text = "tau = T r / J"; prs.save(os.path.join(self.tmp, "b.pptx"))
        self.assertIn("Hooke's law", academic.extract_pages(os.path.join(self.tmp, "a.docx"))[0])
        self.assertEqual(academic.extract_pages(os.path.join(self.tmp, "b.pptx")), ["Torsion\ntau = T r / J"])
        open(os.path.join(self.tmp, "x.exe"), "w").close()
        with self.assertRaises(ValueError):
            academic.add_file(os.path.join(self.tmp, "x.exe"))

    def test_reads_pdf_pages(self):
        try:
            from reportlab.pdfgen import canvas
        except ImportError:
            self.skipTest("reportlab not installed")
        path = os.path.join(self.tmp, "c.pdf")
        c = canvas.Canvas(path); c.drawString(72, 720, "Page one: deflection"); c.showPage()
        c.drawString(72, 720, "Page two: buckling"); c.save()
        entry = academic.add_file(path, "problems")
        self.assertEqual(entry["pages"], 2)
        self.assertEqual(academic.search("buckling")[0]["label"], "c.pdf, p. 2")

    def test_generate_fill_blanks_and_check_answers(self):
        entry = academic.add_file(self.notes, "lecture")
        ai = {"items": [
            {"sentence": "For a simply supported beam with a midspan point load, M_max = ____.", "answer": "P*L/4", "accept": ["PL/4"], "source": "Beams lecture.txt, part 1"},
            {"sentence": "Concrete cover protects steel from corrosion.", "answer": "corrosion", "source": "part 2"},   # no blank: fixed
            {"sentence": "", "answer": "x"},                                                                         # broken: dropped
        ]}
        client = mock.Mock(); client.chat.completions.create.return_value = fake_reply(json.dumps(ai))
        with mock.patch.object(academic, "_client", return_value=client):
            review = academic.generate_review([entry["id"]], "fill_blanks", 5)
        self.assertEqual(len(review["items"]), 2)
        self.assertIn("____", review["items"][1]["sentence"])
        prompt = client.chat.completions.create.call_args.kwargs["messages"][0]["content"]
        self.assertIn("[Beams lecture.txt, part 1]", prompt)                     # the AI sees page labels to cite
        self.assertEqual(client.chat.completions.create.call_args.kwargs["response_format"], {"type": "json_object"})
        item = review["items"][0]
        self.assertTrue(academic.check_blank("p*l/4", item))
        self.assertTrue(academic.check_blank(" PL/4 ", item))
        self.assertFalse(academic.check_blank("P*L/8", item))
        self.assertTrue(academic.check_blank("corrsion", review["items"][1]))   # one small typo is OK
        self.assertEqual(academic.list_reviews()[0]["count"], 2)
        self.assertEqual(academic.save_score(review["id"], 1, 2), 50)
        self.assertEqual(academic.save_score(review["id"], 2, 2), 100)
        self.assertEqual(academic.get_review(review["id"])["best"], 100)

    def test_mcq_validation(self):
        entry = academic.add_file(self.notes, "lecture")
        ai = {"items": [{"question": "Q1", "choices": ["a", "b"], "answer_index": 1},
                        {"question": "Q2", "choices": ["a", "b"], "answer_index": 5}]}       # bad index: dropped
        client = mock.Mock(); client.chat.completions.create.return_value = fake_reply(json.dumps(ai))
        with mock.patch.object(academic, "_client", return_value=client):
            self.assertEqual(len(academic.generate_review([entry["id"]], "mcq")["items"]), 1)
        client.chat.completions.create.return_value = fake_reply(json.dumps({"items": []}))
        with mock.patch.object(academic, "_client", return_value=client), self.assertRaises(ValueError):
            academic.generate_review([entry["id"]], "mcq")

    def test_assignment_help_uses_my_notes(self):
        academic.add_file(self.notes, "lecture")
        client = mock.Mock(); client.chat.completions.create.return_value = fake_reply("Step 1: ...")
        with mock.patch.object(academic, "_client", return_value=client):
            r = academic.assignment_help("Find the maximum moment of a simply supported beam with point load P at midspan", "hints")
        prompt = client.chat.completions.create.call_args.kwargs["messages"][0]["content"]
        self.assertIn("only short hints", prompt)
        self.assertIn("P*L/4", prompt)
        self.assertEqual(r["used"][0], "Beams lecture.txt, part 1")


class CoderTests(unittest.TestCase):
    def setUp(self):
        self.ws = tempfile.mkdtemp()
        with open(os.path.join(self.ws, "calc.py"), "w") as f:
            f.write("def add(a, b):\n    return a - b\n")
        with open(os.path.join(self.ws, ".env"), "w") as f:
            f.write("KEY=secret")
        self.approvals = []
        self.answer = True
        self.agent = coder.CodingAgent(self.ws, approve=lambda k, d: self.approvals.append((k, d)) or self.answer)

    def tearDown(self):
        shutil.rmtree(self.ws)

    def test_read_search_and_safety(self):
        self.assertIn("calc.py", self.agent.list_files())
        self.assertNotIn(".env", self.agent.list_files())
        self.assertIn("2|     return a - b", self.agent.read_file("calc.py"))
        self.assertIn("calc.py:2:", self.agent.search_code("return"))
        self.assertIn("outside the project", self.agent.run_tool("read_file", '{"path": "../../etc/passwd"}'))
        self.assertIn("secrets", self.agent.run_tool("read_file", '{"path": ".env"}'))

    def test_edit_needs_approval(self):
        self.answer = False
        self.assertIn("rejected", self.agent.edit_file("calc.py", "a - b", "a + b"))
        with open(os.path.join(self.ws, "calc.py")) as f:
            self.assertIn("a - b", f.read())
        self.assertIn("-    return a - b", self.approvals[0][1]["diff"])
        self.answer = True
        self.assertEqual(self.agent.edit_file("calc.py", "a - b", "a + b", "fix add"), "Edited calc.py.")
        with open(os.path.join(self.ws, "calc.py")) as f:
            self.assertIn("a + b", f.read())
        self.assertIn("must match exactly once", self.agent.edit_file("calc.py", "nothing like this", "x"))

    def test_run_command_needs_approval(self):
        self.answer = False
        self.assertIn("rejected", self.agent.run_command("echo hi"))
        self.answer = True
        self.assertIn("hi", self.agent.run_command("echo hi"))
        self.assertEqual(self.approvals[-1][0], "run")

    def test_agent_loop(self):
        replies = [fake_reply("", [tool_call("read_file", {"path": "calc.py"}, "c1")]),
                   fake_reply("", [tool_call("edit_file", {"path": "calc.py", "old_text": "a - b", "new_text": "a + b"}, "c2")]),
                   fake_reply("Fixed: add() was subtracting. Run your tests to check.")]
        client = mock.Mock(); client.chat.completions.create.side_effect = replies
        with mock.patch("pinakamain.get_client", return_value=client):
            reply = self.agent.ask("my add function is wrong")
        self.assertIn("Fixed", reply)
        with open(os.path.join(self.ws, "calc.py")) as f:
            self.assertIn("a + b", f.read())
        self.assertEqual(self.agent.history[-1]["role"], "assistant")
        self.assertIn("project folder first", coder.CodingAgent().run_command("dir"))   # nothing runs without a folder


class PlannerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.p = mock.patch.object(f1, "DATA_FILE", os.path.join(self.tmp, "data.json"))
        self.p.start()
        f1.user_data = f1._empty_data()

    def tearDown(self):
        self.p.stop()
        shutil.rmtree(self.tmp)

    def test_day_view_combines_everything(self):
        f1.add_event("Dentist", "2026-10-21", "15:00")
        f1.add_task("Lab report", "10:00", "high", "2026-10-21")
        f1.add_birthday("Mia", "2004-10-21")
        f1.add_reminder("Water plants", "2026-10-19 08:00", "daily")
        f1.add_schedule("Wednesday", "Gym", "18:00")
        day = f1.planner_day("2026-10-21")          # a Wednesday
        self.assertEqual([(i["kind"], i["title"]) for i in day], [
            ("birthday", "Mia's birthday"), ("reminder", "Water plants"), ("task", "Lab report"),
            ("event", "Dentist"), ("weekly", "Gym")])
        self.assertEqual(day[0]["detail"], "turns 22")
        month = f1.planner_month(2026, 10)
        self.assertEqual(month["2026-10-21"], {"birthday": 1, "reminder": 1, "task": 1, "event": 1, "weekly": 1})
        self.assertNotIn("2026-10-18", month)                                  # before the daily reminder started

    def test_repeating_reminders_move_forward(self):
        f1.add_reminder("Standup", "2026-10-23 09:00", "weekdays")              # a Friday
        self.assertEqual(f1.due_alerts(datetime.datetime(2026, 10, 23, 9, 0)), ["Reminder: Standup"])
        self.assertEqual(f1.user_data["reminders"][0]["date_time"], "2026-10-26 09:00")   # skips the weekend
        f1.add_reminder("Rent", "2026-01-31 09:00", "monthly")
        f1.due_alerts(datetime.datetime(2026, 1, 31, 9, 0))
        self.assertEqual(f1.user_data["reminders"][1]["date_time"], "2026-02-28 09:00")   # short month
        f1.add_reminder("Pill", "2026-10-20 08:00", "daily")
        self.assertEqual(f1.due_alerts(datetime.datetime(2026, 10, 25, 12, 0))[-1], "Reminder: Pill")   # PC was off
        self.assertEqual(f1.user_data["reminders"][2]["date_time"], "2026-10-26 08:00")
        self.assertIn("should look like", f1.add_reminder("x", "tomorrow"))

    def test_complete_task(self):
        f1.add_task("Read chapter 3", None)
        self.assertEqual(f1.complete_task("chapter 3"), "Marked 'Read chapter 3' as done.")
        self.assertTrue(f1.user_data["tasks"][0]["done"])
        f1.complete_task("Read chapter 3", False)
        self.assertFalse(f1.user_data["tasks"][0]["done"])


if __name__ == "__main__":
    unittest.main()
